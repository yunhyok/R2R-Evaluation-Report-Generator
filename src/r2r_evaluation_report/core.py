"""Dependency-light parsing, preflight, and evaluation for R2R datasets.

This module intentionally keeps data parsing separate from the Qt user interface and
workbook writer.  It is consequently also the authoritative place for the fixed
26 x 38 input contract.
"""

from __future__ import annotations

import csv
import hashlib
import math
import re
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from statistics import fmean
from typing import Any, Literal

from openpyxl import load_workbook

DatasetKind = Literal["measurement", "prediction"]
Coordinate = tuple[int, int]
RecordKey = tuple[str, int, int]
GRID_ROWS = 26
GRID_NODES = 38
GRID_SIZE = GRID_ROWS * GRID_NODES
TARGET_CLASSES = ("Normal", "Open", "Short")


class DataContractError(ValueError):
    """Raised when an input source does not satisfy the R2R data contract."""


def _normalise(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", unicodedata.normalize("NFKC", str(value)).casefold())


_HEADERS: dict[DatasetKind, dict[str, tuple[str, ...]]] = {
    "measurement": {
        "name": ("name", "sample name", "sample", "device name", "id"),
        "row": ("row", "r"),
        "node": ("node", "column", "col", "n"),
        "status": ("status", "measurement status", "result", "measurement result"),
    },
    "prediction": {
        "name": ("name", "sample name", "sample", "device name", "id"),
        "row": ("row", "r"),
        "node": ("node", "column", "col", "n"),
        "prediction": ("prediction", "predicted class", "predicted", "pred", "class"),
    },
}
_OPTIONAL_HEADERS = {
    "confidence": ("confidence", "confidence score", "score", "probability"),
    "review_required": ("review required", "review_required", "review", "requires review"),
    "prob_Normal": ("prob normal", "prob_normal", "normal probability", "p normal"),
    "prob_Open": ("prob open", "prob_open", "open probability", "p open"),
    "prob_Short": ("prob short", "prob_short", "short probability", "p short"),
    "provenance": ("provenance", "origin", "traceability"),
    "source": ("source", "source file", "input source"),
}
_DEFAULT_STATUS_RULES = {
    "pass": "Normal",
    "noactive": "Normal",
    "none": "Normal",
    "open": "Open",
    "short": "Short",
    "nogateeffect": "Exclude",
}
_DEFAULT_BINARY_STATUS_RULES = {
    "pass": "Pass",
    "noactive": "Fail",
    "none": "Fail",
    "open": "Fail",
    "short": "Fail",
    "nogateeffect": "Fail",
}
_PREDICTION_CLASSES = {_normalise(label): label for label in TARGET_CLASSES}


@dataclass(frozen=True)
class SourceMetadata:
    path: str
    sha256: str
    size: int
    mtime_ns: int
    encoding: str | None = None
    worksheet: str | None = None


@dataclass(frozen=True)
class WorksheetInfo:
    title: str
    header_row: int
    headers: tuple[str, ...]


@dataclass(frozen=True)
class R2RRecord:
    name: str
    row: int
    node: int
    value: str
    confidence: float | None = None
    review_required: bool | None = None
    prob_normal: float | None = None
    prob_open: float | None = None
    prob_short: float | None = None
    provenance: str | None = None
    source: str | None = None
    input_row: int | None = None

    @property
    def coordinate(self) -> Coordinate:
        return (self.row, self.node)

    @property
    def key(self) -> RecordKey:
        return (_normalise(self.name), self.row, self.node)


@dataclass(frozen=True)
class ParsedSheet:
    kind: DatasetKind
    source: SourceMetadata
    records: tuple[R2RRecord, ...]
    header_map: Mapping[str, str]
    sample_name: str | None = None

    @property
    def title(self) -> str:
        return self.sample_name or self.source.worksheet or "CSV"


@dataclass(frozen=True)
class ParsedDataset:
    kind: DatasetKind
    path: str
    sheets: tuple[ParsedSheet, ...]


@dataclass(frozen=True)
class MappingProposal:
    measurement_sheet: str
    prediction_sheet: str
    method: Literal["exact", "signature", "heuristic"]
    requires_confirmation: bool
    signature: tuple[str, ...] = ()


@dataclass(frozen=True)
class MappingAudit:
    proposals: tuple[MappingProposal, ...]
    unmatched_measurement_sheets: tuple[str, ...]
    unmatched_prediction_sheets: tuple[str, ...]
    prediction_only_sheets: tuple[str, ...]


@dataclass(frozen=True)
class JoinedRecord:
    measurement: R2RRecord | None
    prediction: R2RRecord | None
    sheet_name: str

    @property
    def key(self) -> RecordKey | None:
        item = self.measurement or self.prediction
        return item.key if item else None


@dataclass(frozen=True)
class ClassMetrics:
    label: str
    precision: float | None
    recall: float | None
    f1: float | None
    support: int
    specificity: float | None
    false_positive_rate: float | None
    false_negative_rate: float | None
    recall_ci95: tuple[float, float] | None


@dataclass(frozen=True)
class Metrics:
    labels: tuple[str, ...]
    matrix: tuple[tuple[int, ...], ...]
    accuracy: float | None
    accuracy_ci95: tuple[float, float] | None
    balanced_accuracy: float | None
    macro_f1: float | None
    weighted_f1: float | None
    per_class: tuple[ClassMetrics, ...]
    total: int


@dataclass(frozen=True)
class ConfidenceReviewStats:
    confidence_count: int
    confidence_mean: float | None
    review_required_count: int
    review_required_rate: float | None
    probability_count: int


@dataclass(frozen=True)
class YieldStats:
    """Operational pass yields from complete measurement/prediction record pairs."""

    total: int
    measurement_pass_count: int
    measurement_pass_rate: float | None
    measurement_pass_ci95: tuple[float, float] | None
    prediction_pass_count: int
    prediction_pass_rate: float | None
    prediction_pass_ci95: tuple[float, float] | None


@dataclass(frozen=True)
class SheetEvaluation:
    measurement_sheet: str
    prediction_sheet: str
    joined: tuple[JoinedRecord, ...]
    excluded: tuple[JoinedRecord, ...]
    raw_status_by_prediction: Mapping[str, Mapping[str, int]]
    mapped_three_class: Metrics | None
    operational_binary: Metrics | None
    confidence_review: ConfidenceReviewStats
    yield_stats: YieldStats | None = None


@dataclass(frozen=True)
class StatusMappingProfile:
    three_class: Mapping[str, str]
    binary: Mapping[str, str]


@dataclass(frozen=True)
class EvaluationResult:
    blocked: bool
    unresolved_measurement_statuses: tuple[str, ...]
    unknown_predictions: tuple[str, ...]
    mappings: MappingAudit
    sheet_evaluations: tuple[SheetEvaluation, ...] = ()
    overall_three_class: Metrics | None = None
    overall_binary: Metrics | None = None
    measurement_sources: tuple[SourceMetadata, ...] = ()
    prediction_sources: tuple[SourceMetadata, ...] = ()
    status_mapping: StatusMappingProfile | None = None
    overall_yield: YieldStats | None = None
    mapping_errors: tuple[str, ...] = ()


def _source_metadata(
    path: Path, *, encoding: str | None = None, worksheet: str | None = None
) -> SourceMetadata:
    stat = path.stat()
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return SourceMetadata(
        str(path),
        digest.hexdigest(),
        stat.st_size,
        stat.st_mtime_ns,
        encoding,
        worksheet,
    )


def _header_map(headers: Iterable[object], kind: DatasetKind) -> dict[str, str] | None:
    original = [str(h).strip() for h in headers if h is not None and str(h).strip()]
    lookup: dict[str, str] = {}
    for header in original:
        lookup.setdefault(_normalise(header), header)
    result: dict[str, str] = {}
    for field_name, aliases in _HEADERS[kind].items():
        found = next(
            (lookup[_normalise(alias)] for alias in aliases if _normalise(alias) in lookup), None
        )
        if not found:
            return None
        result[field_name] = found
    for field_name, aliases in _OPTIONAL_HEADERS.items():
        found = next(
            (lookup[_normalise(alias)] for alias in aliases if _normalise(alias) in lookup), None
        )
        if found:
            result[field_name] = found
    return result


def _infer_kind(headers: Iterable[object]) -> DatasetKind | None:
    matches = [kind for kind in ("measurement", "prediction") if _header_map(headers, kind)]
    return matches[0] if len(matches) == 1 else None


def discover_worksheets(
    path: str | Path, kind: DatasetKind | None = None
) -> tuple[WorksheetInfo, ...]:
    """Return all XLSX sheets containing a valid header row (first 50 rows).

    CSV has one virtual worksheet named ``CSV``.  Legacy ``.xls`` is rejected
    deliberately because openpyxl does not provide a reliable reader for it.
    """
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".xls":
        raise DataContractError(".xls is unsupported; save the source as .xlsx or .csv")
    if suffix == ".csv":
        headers, _rows, _encoding = _read_csv(source)
        if (kind and not _header_map(headers, kind)) or (not kind and not _infer_kind(headers)):
            return ()
        return (WorksheetInfo("CSV", 1, tuple(headers)),)
    if suffix != ".xlsx":
        raise DataContractError("only .csv and .xlsx sources are supported")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        found: list[WorksheetInfo] = []
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1
            ):
                is_contract_header = _header_map(row, kind) if kind else _infer_kind(row)
                if is_contract_header:
                    found.append(
                        WorksheetInfo(
                            worksheet.title,
                            row_number,
                            tuple("" if value is None else str(value) for value in row),
                        )
                    )
                    break
        return tuple(found)
    finally:
        workbook.close()


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]], str]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise DataContractError("CSV has no header row")
                return list(reader.fieldnames), list(reader), encoding
        except UnicodeDecodeError as error:
            last_error = error
    raise DataContractError(f"cannot decode CSV with UTF-8-SIG, UTF-8, or CP949: {last_error}")


def _int_coordinate(value: Any, field: str, context: str) -> int:
    if isinstance(value, bool):
        raise DataContractError(f"{context}: {field} must be an integer")
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            raise DataContractError(f"{context}: {field} must be an integer, got {value!r}")
        return int(value)
    text = str(value).strip()
    if not re.fullmatch(r"[+-]?\d+", text):
        raise DataContractError(f"{context}: {field} must be an integer, got {value!r}")
    return int(text)


def _optional_float(value: Any, field: str, context: str) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise DataContractError(f"{context}: {field} must be numeric") from error
    if not math.isfinite(number):
        raise DataContractError(f"{context}: {field} must be finite")
    return number


def _optional_bool(value: Any, context: str) -> bool | None:
    if value is None or str(value).strip() == "":
        return None
    normal = _normalise(value)
    if normal in {"true", "yes", "y", "1"}:
        return True
    if normal in {"false", "no", "n", "0"}:
        return False
    raise DataContractError(f"{context}: review_required must be true/false")


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _parse_rows(
    rows: Iterable[tuple[int, Mapping[str, Any]]],
    header_map: Mapping[str, str],
    kind: DatasetKind,
    metadata: SourceMetadata,
) -> tuple[R2RRecord, ...]:
    records: list[R2RRecord] = []
    seen_by_name: dict[str, set[Coordinate]] = defaultdict(set)
    for row_number, source_row in rows:
        # Completely blank tail rows are not records; partially blank rows remain errors.
        if not any(value is not None and str(value).strip() for value in source_row.values()):
            continue
        context = f"{metadata.path} [{metadata.worksheet or 'CSV'}] row {row_number}"
        name = _as_text(source_row.get(header_map["name"]))
        value = _as_text(
            source_row.get(header_map["status" if kind == "measurement" else "prediction"])
        )
        if not name or not value:
            value_name = "Status" if kind == "measurement" else "prediction"
            raise DataContractError(f"{context}: Name and {value_name} are required")
        row = _int_coordinate(source_row.get(header_map["row"]), "Row", context)
        node = _int_coordinate(source_row.get(header_map["node"]), "Node", context)
        if not 1 <= row <= GRID_ROWS or not 1 <= node <= GRID_NODES:
            raise DataContractError(
                f"{context}: coordinate ({row}, {node}) is outside Row 1..26 / Node 1..38"
            )
        coordinate = (row, node)
        normalised_name = _normalise(name)
        if coordinate in seen_by_name[normalised_name]:
            raise DataContractError(
                f"{context}: duplicate coordinate ({row}, {node}) for Name {name!r}"
            )
        seen_by_name[normalised_name].add(coordinate)
        confidence = (
            _optional_float(source_row.get(header_map.get("confidence")), "confidence", context)
            if "confidence" in header_map
            else None
        )
        if confidence is not None and not 0 <= confidence <= 1:
            raise DataContractError(f"{context}: confidence must be in [0, 1]")
        probabilities = tuple(
            _optional_float(source_row.get(header_map.get(field_name)), field_name, context)
            if field_name in header_map
            else None
            for field_name in ("prob_Normal", "prob_Open", "prob_Short")
        )
        if any(item is not None for item in probabilities):
            if any(item is None for item in probabilities):
                raise DataContractError(
                    f"{context}: prob_Normal, prob_Open, and prob_Short must be supplied together"
                )
            if any(not 0 <= item <= 1 for item in probabilities if item is not None):
                raise DataContractError(f"{context}: probabilities must be in [0, 1]")
            if abs(sum(item for item in probabilities if item is not None) - 1.0) > 1e-5:
                raise DataContractError(f"{context}: probabilities must sum to 1 within 1e-5")
        records.append(
            R2RRecord(
                name,
                row,
                node,
                value,
                confidence,
                _optional_bool(source_row.get(header_map["review_required"]), context)
                if "review_required" in header_map
                else None,
                *probabilities,
                _as_text(source_row.get(header_map["provenance"]))
                if "provenance" in header_map
                else None,
                _as_text(source_row.get(header_map["source"])) if "source" in header_map else None,
                row_number,
            )
        )
    if not seen_by_name:
        location = f"{metadata.path} [{metadata.worksheet or 'CSV'}]"
        raise DataContractError(
            f"{location}: expected at least one Name with exactly {GRID_SIZE} unique coordinates"
        )
    for normalised_name, coordinates in seen_by_name.items():
        if len(coordinates) != GRID_SIZE:
            missing = GRID_SIZE - len(coordinates)
            location = f"{metadata.path} [{metadata.worksheet or 'CSV'}]"
            raise DataContractError(
                f"{location}: Name {normalised_name!r} expected exactly {GRID_SIZE} "
                f"unique coordinates; found {len(coordinates)} ({missing:+d} vs expected)"
            )
    return tuple(records)


def _sample_sheets(
    kind: DatasetKind,
    metadata: SourceMetadata,
    records: tuple[R2RRecord, ...],
    header_map: Mapping[str, str],
) -> tuple[ParsedSheet, ...]:
    """Split a header-bearing input worksheet into its complete Name sample units."""
    grouped: dict[str, list[R2RRecord]] = defaultdict(list)
    display_names: dict[str, str] = {}
    for record in records:
        normalised_name = _normalise(record.name)
        grouped[normalised_name].append(record)
        display_names.setdefault(normalised_name, record.name)
    return tuple(
        ParsedSheet(kind, metadata, tuple(group), header_map, display_names[normalised_name])
        for normalised_name, group in grouped.items()
    )


def _ensure_unique_samples(sheets: Sequence[ParsedSheet]) -> None:
    """Reject a sample Name repeated in separate source worksheets."""
    source_by_name: dict[str, str | None] = {}
    for sheet in sheets:
        normalised_name = _normalise(sheet.title)
        previous_worksheet = source_by_name.get(normalised_name)
        if previous_worksheet is not None and previous_worksheet != sheet.source.worksheet:
            raise DataContractError(
                f"duplicate equivalent Name {sheet.title!r} across worksheets "
                f"{previous_worksheet!r} and {sheet.source.worksheet!r}"
            )
        source_by_name[normalised_name] = sheet.source.worksheet


def parse_dataset(
    path: str | Path,
    kind: DatasetKind,
    worksheets: str | Sequence[str] | None = None,
) -> ParsedDataset:
    """Parse every selected contract-bearing worksheet into immutable records."""
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".xls":
        raise DataContractError(".xls is unsupported; save the source as .xlsx or .csv")
    if suffix == ".csv":
        if worksheets not in (None, "CSV", ("CSV",)):
            raise DataContractError("CSV exposes only the virtual worksheet 'CSV'")
        headers, rows, encoding = _read_csv(source)
        header_map = _header_map(headers, kind)
        if not header_map:
            raise DataContractError(f"CSV is missing required {kind} headers")
        metadata = _source_metadata(source, encoding=encoding, worksheet="CSV")
        records = _parse_rows(enumerate(rows, start=2), header_map, kind, metadata)
        parsed = _sample_sheets(kind, metadata, records, header_map)
        return ParsedDataset(kind, str(source), parsed)
    if suffix != ".xlsx":
        raise DataContractError("only .csv and .xlsx sources are supported")
    candidates = discover_worksheets(source, kind)
    names = {worksheets} if isinstance(worksheets, str) else set(worksheets or ())
    selected = tuple(info for info in candidates if not names or info.title in names)
    missing_names = names - {item.title for item in candidates}
    if missing_names:
        raise DataContractError(
            f"selected worksheet(s) have no valid {kind} header: {sorted(missing_names)}"
        )
    if not selected:
        raise DataContractError(f"no {kind} worksheets were discovered")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        parsed: list[ParsedSheet] = []
        base_metadata = _source_metadata(source)
        for info in selected:
            worksheet = workbook[info.title]
            header_values = next(
                worksheet.iter_rows(
                    min_row=info.header_row, max_row=info.header_row, values_only=True
                )
            )
            header_map = _header_map(header_values, kind)
            assert header_map is not None
            metadata = SourceMetadata(
                base_metadata.path,
                base_metadata.sha256,
                base_metadata.size,
                base_metadata.mtime_ns,
                worksheet=info.title,
            )
            rows = (
                (
                    row_number,
                    {
                        str(header): value
                        for header, value in zip(header_values, values, strict=False)
                        if header is not None
                    },
                )
                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=info.header_row + 1, values_only=True),
                    start=info.header_row + 1,
                )
            )
            records = _parse_rows(rows, header_map, kind, metadata)
            parsed.extend(_sample_sheets(kind, metadata, records, header_map))
        _ensure_unique_samples(parsed)
        return ParsedDataset(kind, str(source), tuple(parsed))
    finally:
        workbook.close()


def _sheet_name(value: ParsedSheet | str) -> str:
    return value.title if isinstance(value, ParsedSheet) else str(value)


def _signature(title: str) -> tuple[str, ...]:
    normal = unicodedata.normalize("NFKC", title).casefold()
    date = re.search(r"(?<!\d)(?:(20\d{2})|(\d{2}))(\d{2})(\d{2})(?!\d)", normal)
    # Underscores are word characters to the regex engine but are separators in
    # the real sample names (for example ``_7kgf_100_sam1``).  Use explicit
    # alphanumeric boundaries so CSV Name values and human-readable titles share
    # the same token contract.
    kg = re.search(r"(?<![a-z0-9])(\d+(?:\.\d+)?)\s*kgf?(?![a-z0-9])", normal)
    sam = re.search(r"(?<![a-z0-9])sam\s*[-_ ]?0*(\d+)(?!\d)", normal)
    values: list[str] = []
    if date:
        year = date.group(1)[2:] if date.group(1) else date.group(2)
        values.append(f"date:{year}{date.group(3)}{date.group(4)}")
    if kg:
        values.append(f"kg:{float(kg.group(1)):g}")
    if sam:
        values.append(f"sam:{int(sam.group(1))}")
    return tuple(values)


def _mapping_signature(sheet: ParsedSheet) -> tuple[str, ...]:
    return _signature(f"{sheet.title} {sheet.source.worksheet or ''}")


def _signature_matches(left: tuple[str, ...], right: tuple[str, ...]) -> bool:
    left_parts = dict(part.split(":", maxsplit=1) for part in left)
    right_parts = dict(part.split(":", maxsplit=1) for part in right)
    if not {"date", "sam"} <= left_parts.keys() & right_parts.keys():
        return False
    if left_parts["date"] != right_parts["date"] or left_parts["sam"] != right_parts["sam"]:
        return False
    return not (
        "kg" in left_parts and "kg" in right_parts and left_parts["kg"] != right_parts["kg"]
    )


def propose_mappings(
    measurement_sheets: Sequence[ParsedSheet], prediction_sheets: Sequence[ParsedSheet]
) -> MappingAudit:
    """Pair sheets deterministically; heuristic suggestions are never auto-approved."""
    remaining_measurements = list(measurement_sheets)
    remaining_predictions = list(prediction_sheets)
    proposals: list[MappingProposal] = []
    for measurement in tuple(remaining_measurements):
        matches = [
            prediction
            for prediction in remaining_predictions
            if _normalise(prediction.title) == _normalise(measurement.title)
        ]
        if len(matches) == 1:
            prediction = matches[0]
            proposals.append(MappingProposal(measurement.title, prediction.title, "exact", False))
            remaining_measurements.remove(measurement)
            remaining_predictions.remove(prediction)
    signature_candidates: list[tuple[ParsedSheet, ParsedSheet, tuple[str, ...]]] = []
    for measurement in tuple(remaining_measurements):
        signature = _mapping_signature(measurement)
        matches = [
            prediction
            for prediction in remaining_predictions
            if _signature_matches(signature, _mapping_signature(prediction))
        ]
        if len(matches) == 1:
            signature_candidates.append((measurement, matches[0], signature))
    prediction_candidate_counts = Counter(
        prediction.title for _measurement, prediction, _signature_value in signature_candidates
    )
    reserved_measurements: set[str] = set()
    reserved_predictions: set[str] = set()
    for measurement, prediction, signature in signature_candidates:
        # A candidate is only safe to present independently when the reverse side
        # is unique as well.  This prevents two checkboxes from reserving the same
        # prediction sample.
        if prediction_candidate_counts[prediction.title] == 1:
            proposals.append(
                MappingProposal(measurement.title, prediction.title, "signature", True, signature)
            )
            reserved_measurements.add(measurement.title)
            reserved_predictions.add(prediction.title)
    # A deterministic best-name candidate is useful UI guidance but is deliberately not an approval.
    heuristic_scores = sorted(
        (
            (
                SequenceMatcher(
                    None, _normalise(measurement.title), _normalise(prediction.title)
                ).ratio(),
                measurement,
                prediction,
            )
            for measurement in remaining_measurements
            if measurement.title not in reserved_measurements
            for prediction in remaining_predictions
            if prediction.title not in reserved_predictions
        ),
        key=lambda item: (
            -item[0],
            _normalise(item[1].title),
            _normalise(item[2].title),
        ),
    )
    for score, measurement, prediction in heuristic_scores:
        if score < 0.45:
            break
        if (
            measurement.title not in reserved_measurements
            and prediction.title not in reserved_predictions
        ):
            proposals.append(
                MappingProposal(measurement.title, prediction.title, "heuristic", True)
            )
            reserved_measurements.add(measurement.title)
            reserved_predictions.add(prediction.title)
    auto_approved_measurements = {
        item.measurement_sheet for item in proposals if not item.requires_confirmation
    }
    auto_approved_predictions = {
        item.prediction_sheet for item in proposals if not item.requires_confirmation
    }
    suggested_predictions = {item.prediction_sheet for item in proposals}
    return MappingAudit(
        tuple(proposals),
        tuple(
            item.title
            for item in measurement_sheets
            if item.title not in auto_approved_measurements
        ),
        tuple(
            item.title for item in prediction_sheets if item.title not in auto_approved_predictions
        ),
        tuple(item.title for item in prediction_sheets if item.title not in suggested_predictions),
    )


def preflight(measurements: ParsedDataset, predictions: ParsedDataset) -> MappingAudit:
    if measurements.kind != "measurement" or predictions.kind != "prediction":
        raise DataContractError("preflight requires a measurement dataset and a prediction dataset")
    return propose_mappings(measurements.sheets, predictions.sheets)


def join_records(measurement: ParsedSheet, prediction: ParsedSheet) -> tuple[JoinedRecord, ...]:
    # Sheet mapping already establishes sample identity.  The raw ``Name`` values
    # intentionally differ between measurement and prediction sources, so joining
    # on the full source key would turn every valid mapping into 1,976 one-sided
    # rows.  Coordinates are the canonical identity within an approved sample pair.
    measured = {record.coordinate: record for record in measurement.records}
    predicted = {record.coordinate: record for record in prediction.records}
    keys = sorted(set(measured) | set(predicted))
    return tuple(
        JoinedRecord(measured.get(key), predicted.get(key), measurement.title) for key in keys
    )


def _wilson(successes: int, total: int) -> tuple[float, float] | None:
    if total == 0:
        return None
    z = 1.959963984540054
    proportion = successes / total
    denominator = 1 + z * z / total
    centre = (proportion + z * z / (2 * total)) / denominator
    spread = (
        z * math.sqrt((proportion * (1 - proportion) + z * z / (4 * total)) / total) / denominator
    )
    return (max(0.0, centre - spread), min(1.0, centre + spread))


def _safe_divide(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def _metrics(
    labels: Sequence[str],
    actual_predicted: Sequence[tuple[str, str]],
    *,
    strict_macro: bool = False,
) -> Metrics:
    labels = tuple(labels)
    index = {label: position for position, label in enumerate(labels)}
    cells = [[0 for _ in labels] for _ in labels]
    for actual, predicted in actual_predicted:
        cells[index[actual]][index[predicted]] += 1
    total = len(actual_predicted)
    correct = sum(cells[position][position] for position in range(len(labels)))
    class_metrics: list[ClassMetrics] = []
    f1s: list[float] = []
    weighted_numerator = 0.0
    recalls: list[float] = []
    supports: list[int] = []
    for position, label in enumerate(labels):
        tp = cells[position][position]
        fn = sum(cells[position]) - tp
        fp = sum(row[position] for row in cells) - tp
        tn = total - tp - fn - fp
        support = tp + fn
        precision = _safe_divide(tp, tp + fp)
        recall = _safe_divide(tp, support)
        f1 = None if support == 0 else _safe_divide(2 * tp, 2 * tp + fp + fn)
        specificity = _safe_divide(tn, tn + fp)
        class_metrics.append(
            ClassMetrics(
                label,
                precision,
                recall,
                f1,
                support,
                specificity,
                None if specificity is None else 1 - specificity,
                _safe_divide(fn, fn + tp),
                _wilson(tp, support),
            )
        )
        if recall is not None:
            recalls.append(recall)
        if f1 is not None:
            f1s.append(f1)
            weighted_numerator += f1 * support
        supports.append(support)
    macro = (
        None
        if strict_macro and any(support == 0 for support in supports)
        else (fmean(f1s) if f1s else None)
    )
    return Metrics(
        labels,
        tuple(tuple(row) for row in cells),
        _safe_divide(correct, total),
        _wilson(correct, total),
        fmean(recalls) if recalls else None,
        macro,
        weighted_numerator / total if total else None,
        tuple(class_metrics),
        total,
    )


def _raw_table(rows: Sequence[JoinedRecord]) -> dict[str, dict[str, int]]:
    table: dict[str, Counter[str]] = defaultdict(Counter)
    for joined in rows:
        if joined.measurement and joined.prediction:
            table[joined.measurement.value][joined.prediction.value] += 1
    return {
        status: dict(sorted(predictions.items())) for status, predictions in sorted(table.items())
    }


def _yield_stats(rows: Sequence[JoinedRecord]) -> YieldStats:
    """Calculate operational yields from rows eligible for evaluation."""
    eligible = [row for row in rows if row.measurement and row.prediction]
    measurement_passes = sum(
        _normalise(row.measurement.value) == "pass" for row in eligible if row.measurement
    )
    prediction_passes = sum(
        _normalise(row.prediction.value) == "normal" for row in eligible if row.prediction
    )
    total = len(eligible)
    return YieldStats(
        total,
        measurement_passes,
        _safe_divide(measurement_passes, total),
        _wilson(measurement_passes, total),
        prediction_passes,
        _safe_divide(prediction_passes, total),
        _wilson(prediction_passes, total),
    )


def _confidence_stats(rows: Sequence[JoinedRecord]) -> ConfidenceReviewStats:
    predictions = [row.prediction for row in rows if row.prediction]
    confidence = [item.confidence for item in predictions if item and item.confidence is not None]
    required = sum(bool(item.review_required) for item in predictions if item)
    probability_count = sum(item.prob_normal is not None for item in predictions if item)
    return ConfidenceReviewStats(
        len(confidence),
        fmean(confidence) if confidence else None,
        required,
        required / len(predictions) if predictions else None,
        probability_count,
    )


def _mapping_for(status: str, rules: Mapping[str, str]) -> str | None:
    return rules.get(_normalise(status))


def evaluate(
    measurements: ParsedDataset,
    predictions: ParsedDataset,
    mappings: Sequence[MappingProposal] | None = None,
    *,
    status_rules: Mapping[str, str] | None = None,
    binary_status_rules: Mapping[str, str] | None = None,
) -> EvaluationResult:
    """Evaluate approved mappings without silently scoring unknown labels.

    Unknown measurement statuses and predictions are reported in the result and
    make it ``blocked``.  A non-default measurement status needs explicit
    three-class and binary rules; prediction labels are fixed to Normal/Open/Short.
    """
    audit = preflight(measurements, predictions)
    approved = (
        tuple(mappings)
        if mappings is not None
        else tuple(item for item in audit.proposals if not item.requires_confirmation)
    )
    measurement_by_title = {sheet.title: sheet for sheet in measurements.sheets}
    prediction_by_title = {sheet.title: sheet for sheet in predictions.sheets}
    rules = dict(_DEFAULT_STATUS_RULES)
    binary_rules = dict(_DEFAULT_BINARY_STATUS_RULES)
    if status_rules:
        allowed_targets = {_normalise(value): value for value in (*TARGET_CLASSES, "Exclude")}
        resolved_rules: dict[str, str] = {}
        for key, value in status_rules.items():
            target = allowed_targets.get(_normalise(value))
            if target is None:
                raise DataContractError(
                    "status rule targets must be Normal, Open, Short, or Exclude"
                )
            resolved_rules[_normalise(key)] = target
        rules.update(resolved_rules)
    if binary_status_rules:
        resolved_binary_rules: dict[str, str] = {}
        for key, value in binary_status_rules.items():
            target = {"pass": "Pass", "fail": "Fail"}.get(_normalise(value))
            if target is None:
                raise DataContractError("binary status rule targets must be Pass or Fail")
            resolved_binary_rules[_normalise(key)] = target
        binary_rules.update(resolved_binary_rules)
    profile = StatusMappingProfile(dict(rules), dict(binary_rules))
    joined_by_sheet: list[tuple[MappingProposal, tuple[JoinedRecord, ...]]] = []
    # Validate selected sources before pair-specific scoring.  This prevents an
    # unknown label from being hidden simply because its sheet is unmatched.
    unresolved = {
        record.value
        for sheet in measurements.sheets
        for record in sheet.records
        if (
            _mapping_for(record.value, rules) is None
            or _mapping_for(record.value, binary_rules) is None
        )
    }
    unknown_predictions = {
        record.value
        for sheet in predictions.sheets
        for record in sheet.records
        if _normalise(record.value) not in _PREDICTION_CLASSES
    }
    mapping_errors: list[str] = []
    paired_measurements: set[str] = set()
    paired_predictions: set[str] = set()
    for proposal in approved:
        if proposal.requires_confirmation:
            mapping_name = f"{proposal.measurement_sheet!r} -> {proposal.prediction_sheet!r}"
            mapping_errors.append(f"{mapping_name} requires confirmation")
            continue
        measurement = measurement_by_title.get(proposal.measurement_sheet)
        prediction = prediction_by_title.get(proposal.prediction_sheet)
        if not measurement or not prediction:
            mapping_errors.append("a selected mapping references an unknown sample")
            continue
        if measurement.title in paired_measurements or prediction.title in paired_predictions:
            mapping_errors.append("approved mappings must be one-to-one")
            continue
        paired_measurements.add(measurement.title)
        paired_predictions.add(prediction.title)
        joined = join_records(measurement, prediction)
        joined_by_sheet.append((proposal, joined))
    missing_measurements = set(measurement_by_title) - paired_measurements
    unpaired_predictions = set(prediction_by_title) - paired_predictions
    final_audit = MappingAudit(
        tuple(proposal for proposal, _joined in joined_by_sheet),
        tuple(sorted(missing_measurements)),
        tuple(sorted(unpaired_predictions)),
        tuple(sorted(unpaired_predictions)),
    )
    if missing_measurements:
        mapping_errors.append(
            f"unmatched measurement samples: {', '.join(sorted(missing_measurements))}"
        )
    if unresolved or unknown_predictions or mapping_errors:
        return EvaluationResult(
            True,
            tuple(sorted(unresolved)),
            tuple(sorted(unknown_predictions)),
            final_audit,
            measurement_sources=tuple(sheet.source for sheet in measurements.sheets),
            prediction_sources=tuple(sheet.source for sheet in predictions.sheets),
            status_mapping=profile,
            mapping_errors=tuple(mapping_errors),
        )
    evaluations: list[SheetEvaluation] = []
    all_three: list[tuple[str, str]] = []
    all_binary: list[tuple[str, str]] = []
    all_complete: list[JoinedRecord] = []
    for proposal, joined in joined_by_sheet:
        included: list[JoinedRecord] = []
        excluded: list[JoinedRecord] = []
        three: list[tuple[str, str]] = []
        binary: list[tuple[str, str]] = []
        for item in joined:
            if not item.measurement or not item.prediction:
                excluded.append(item)
                continue
            mapped = _mapping_for(item.measurement.value, rules)
            binary_actual = _mapping_for(item.measurement.value, binary_rules)
            assert mapped is not None
            assert binary_actual is not None
            predicted = _PREDICTION_CLASSES[_normalise(item.prediction.value)]
            binary.append((binary_actual, "Pass" if predicted == "Normal" else "Fail"))
            all_complete.append(item)
            if mapped == "Exclude":
                excluded.append(item)
                continue
            included.append(item)
            three.append((mapped, predicted))
        all_three.extend(three)
        all_binary.extend(binary)
        evaluations.append(
            SheetEvaluation(
                proposal.measurement_sheet,
                proposal.prediction_sheet,
                tuple(joined),
                tuple(excluded),
                _raw_table(joined),
                _metrics(TARGET_CLASSES, three, strict_macro=True),
                _metrics(("Pass", "Fail"), binary),
                _confidence_stats(joined),
                _yield_stats(joined),
            )
        )
    return EvaluationResult(
        False,
        (),
        (),
        final_audit,
        tuple(evaluations),
        _metrics(TARGET_CLASSES, all_three, strict_macro=True),
        _metrics(("Pass", "Fail"), all_binary),
        tuple(sheet.source for sheet in measurements.sheets),
        tuple(sheet.source for sheet in predictions.sheets),
        profile,
        _yield_stats(all_complete),
    )


def build_synthetic_evaluation() -> EvaluationResult:
    """Build a deterministic, contract-complete evaluation for installer smoke tests.

    It never reads a user file.  The generated records include every default
    measurement rule, all three prediction classes, a controlled error pattern,
    confidence/probability fields, and auditable ``No Gate Effect`` exclusions.
    """
    digest = hashlib.sha256(b"r2r-evaluation-report synthetic v1").hexdigest()
    title = "Synthetic SAM 11"
    measurement_source = SourceMetadata("<synthetic-measurement>", digest, 0, 0, "synthetic", title)
    prediction_source = SourceMetadata("<synthetic-prediction>", digest, 0, 0, "synthetic", title)
    statuses = ("Pass", "No Active", "None", "Open", "Short", "No Gate Effect")
    measurements: list[R2RRecord] = []
    predictions: list[R2RRecord] = []
    for index, (row, node) in enumerate(
        (row, node) for row in range(1, 27) for node in range(1, 39)
    ):
        status = statuses[index % len(statuses)]
        mapped = _DEFAULT_STATUS_RULES[_normalise(status)]
        predicted = "Normal" if mapped in {"Normal", "Exclude"} else mapped
        if index % 17 == 0:
            predicted = "Open" if predicted == "Normal" else "Normal"
        probabilities = {
            "Normal": (0.9, 0.05, 0.05),
            "Open": (0.05, 0.9, 0.05),
            "Short": (0.05, 0.05, 0.9),
        }[predicted]
        measurements.append(R2RRecord("Synthetic", row, node, status, input_row=index + 2))
        predictions.append(
            R2RRecord(
                "Synthetic",
                row,
                node,
                predicted,
                0.9,
                index % 23 == 0,
                *probabilities,
                "synthetic-v1",
                "generated",
                index + 2,
            )
        )
    measured = ParsedDataset(
        "measurement",
        measurement_source.path,
        (ParsedSheet("measurement", measurement_source, tuple(measurements), {}),),
    )
    predicted = ParsedDataset(
        "prediction",
        prediction_source.path,
        (ParsedSheet("prediction", prediction_source, tuple(predictions), {}),),
    )
    return evaluate(measured, predicted)
