"""Qt-free backend for the linear wizard: inspect → load/align → generate.

Everything here is plain Python so the wizard's decisions can be tested
without a display and reused by the CLI ``--profile`` path.
"""

from __future__ import annotations

from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from threading import Event
from typing import Any

from . import profile as profile_module
from .core import (
    DataContractError,
    ParsedDataset,
    SampleCandidate,
    _read_csv,
    discover_worksheets,
    inspect_samples,
    inspector_model_columns,
    selection_needed,
)
from .profile import (
    AlignmentProposal,
    AlignmentSpec,
    ComparisonSpec,
    DatasetSpec,
    ProfileResult,
    ReportProfile,
)
from .schemes import EXCLUDE, SchemeRegistry, load_registry

ProgressCallback = Callable[[int, str], None]


class GenerationCancelled(RuntimeError):
    """Raised when the cancel event is set between stages."""


@dataclass(frozen=True)
class SourceInfo:
    path: str
    worksheets: tuple[str, ...]
    inspector_models: tuple[str, ...]

    @property
    def is_inspector(self) -> bool:
        return bool(self.inspector_models)


def _long_format_models(headers: Sequence[str], rows) -> list[str]:
    """Distinct ``provider:model_id`` values of a long-format Inspector export."""
    provider = next((h for h in headers if h and h.strip().casefold() == "provider"), None)
    model = next((h for h in headers if h and h.strip().casefold() == "model_id"), None)
    if not provider or not model:
        return []
    found: list[str] = []
    for row in rows:
        key = f"{row.get(provider) or ''}:{row.get(model) or ''}"
        if key != ":" and key not in found:
            found.append(key)
    return found


def inspect_source(path: str | Path) -> SourceInfo:
    """Worksheets and Inspector model identifiers (``provider:model_id``) in ``path``.

    Matrix-style exports name the models in their headers; long-format exports
    (``verdict`` column) are scanned for distinct ``provider:model_id`` pairs, so
    the wizard always offers concrete choices.
    """
    source = Path(path)
    if source.suffix.casefold() == ".csv":
        headers, rows, _encoding = _read_csv(source)
        models = list(inspector_model_columns(headers))
        if models == ["verdict"]:
            models = _long_format_models(headers, rows)
        return SourceInfo(str(source), ("CSV",), tuple(models))
    infos = discover_worksheets(source, "label")
    models: list[str] = []
    long_sheets: list[str] = []
    for info in infos:
        columns = inspector_model_columns(info.headers)
        if columns == ("verdict",):
            long_sheets.append(info.title)
            continue
        for model in columns:
            if model not in models:
                models.append(model)
    if not models and long_sheets:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            worksheet = workbook[long_sheets[0]]
            iterator = worksheet.iter_rows(values_only=True)
            headers = [str(h) if h is not None else "" for h in next(iterator)]
            rows = (dict(zip(headers, values, strict=False)) for values in iterator)
            models = _long_format_models(headers, rows)
        finally:
            workbook.close()
    return SourceInfo(str(source), tuple(info.title for info in infos), tuple(models))


def default_title(path: str | Path, model_column: str | None = None) -> str:
    """Dataset title shown before the operator edits it.

    Inspector exports hold one column per model, and one report may compare
    several of them from the same file, so the model is part of the title.
    """
    stem = Path(path).stem
    return f"{stem} [{model_column}]" if model_column else stem


def inspect_dataset_samples(spec: DatasetSpec) -> tuple[SampleCandidate, ...]:
    """Every sample block of one dataset (repeat Names and incomplete arrays included)."""
    return inspect_samples(
        spec.path,
        "label",
        spec.worksheets or None,
        grid=spec.grid,
        model_column=spec.model_column,
    )


def pending_sample_selections(
    specs: Sequence[DatasetSpec], progress: ProgressCallback | None = None
) -> dict[str, tuple[SampleCandidate, ...]]:
    """Datasets whose blocks need an operator decision and carry no saved selection."""
    pending: dict[str, tuple[SampleCandidate, ...]] = {}
    for index, spec in enumerate(specs):
        if spec.samples:
            continue
        if progress:
            progress(int(5 + 90 * index / max(len(specs), 1)), f"{spec.label} 샘플 블록 검사 중…")
        candidates = inspect_dataset_samples(spec)
        if selection_needed(candidates):
            pending[spec.id] = candidates
    if progress:
        progress(100, "완료")
    return pending


def selection_summary(candidates: Sequence[SampleCandidate]) -> dict[str, list[SampleCandidate]]:
    """Candidates grouped by Name, in first-seen order."""
    grouped: dict[str, list[SampleCandidate]] = {}
    for item in candidates:
        grouped.setdefault(item.name, []).append(item)
    return grouped


def validate_selection(
    candidates: Sequence[SampleCandidate], selection: Mapping[str, str]
) -> list[str]:
    problems: list[str] = []
    by_key = {item.key: item for item in candidates}
    titles: dict[str, str] = {}
    for key, title in selection.items():
        item = by_key.get(key)
        if item is None:
            problems.append(f"알 수 없는 블록: {key}")
            continue
        if not item.complete:
            problems.append(
                f"{item.default_title}: 좌표 {item.unique_coordinates}/{item.grid_size}개 — "
                "불완전한 블록은 포함할 수 없습니다."
            )
        normalised = "".join(ch for ch in title.casefold() if ch.isalnum())
        if not normalised:
            problems.append(f"{item.default_title}: 샘플 이름이 비어 있습니다.")
        elif normalised in titles:
            problems.append(f"샘플 이름 {title!r}이(가) {titles[normalised]!r}와 중복됩니다.")
        else:
            titles[normalised] = item.default_title
    if not selection:
        problems.append("포함할 샘플이 없습니다.")
    return problems


@dataclass(frozen=True)
class LoadResult:
    datasets: Mapping[str, ParsedDataset]
    proposals: tuple[AlignmentProposal, ...]
    suggested_schemes: Mapping[str, str | None]
    raw_labels: Mapping[str, tuple[str, ...]]


def load_and_propose(
    specs: Sequence[DatasetSpec],
    registry: SchemeRegistry | None = None,
    progress: ProgressCallback | None = None,
) -> LoadResult:
    """Parse every dataset, suggest a scheme per dataset and propose sample alignments."""
    registry = registry or load_registry()
    if not specs:
        raise DataContractError("add at least one dataset")
    profile = ReportProfile(tuple(specs))
    datasets: dict[str, ParsedDataset] = {}
    for index, spec in enumerate(specs):
        if progress:
            progress(int(5 + 60 * index / len(specs)), f"{spec.label} 파일을 읽는 중…")
        datasets[spec.id] = profile_module.load_datasets(ReportProfile((spec,)))[spec.id]
    if progress:
        progress(75, "샘플 정렬 후보를 찾는 중…")
    proposals = profile_module.propose_alignments(profile, datasets)
    suggestions: dict[str, str | None] = {}
    raw_labels: dict[str, tuple[str, ...]] = {}
    for spec in specs:
        labels = datasets[spec.id].raw_labels
        raw_labels[spec.id] = labels
        if spec.scheme and registry.scheme(spec.scheme):
            suggestions[spec.id] = spec.scheme
            continue
        matches = registry.identify(labels)
        suggestions[spec.id] = matches[0] if matches else None
    if progress:
        progress(100, "완료")
    return LoadResult(datasets, proposals, suggestions, raw_labels)


def alignment_rows(
    specs: Sequence[DatasetSpec], datasets: Mapping[str, ParsedDataset], load: LoadResult
) -> list[dict[str, Any]]:
    """One row per primary sample: suggested member per dataset and its confirmation need."""
    primary = specs[0]
    rows: list[dict[str, Any]] = []
    by_dataset = {item.dataset_id: item for item in load.proposals}
    for sheet in datasets[primary.id].sheets:
        row: dict[str, Any] = {"sample": sheet.title, "members": {}, "auto": {}}
        for spec in specs[1:]:
            proposal = next(
                (
                    item
                    for item in by_dataset[spec.id].proposals
                    if item.measurement_sheet == sheet.title
                ),
                None,
            )
            row["members"][spec.id] = proposal.prediction_sheet if proposal else None
            row["auto"][spec.id] = bool(proposal and not proposal.requires_confirmation)
        rows.append(row)
    return rows


def member_choices(spec: DatasetSpec, datasets: Mapping[str, ParsedDataset]) -> tuple[str, ...]:
    return tuple(sheet.title for sheet in datasets[spec.id].sheets)


def build_alignments(
    profile: ReportProfile,
    confirmed: Iterable[tuple[str, str, str]],
    primary_samples: Iterable[str] = (),
) -> tuple[AlignmentSpec, ...]:
    return profile_module.alignments_from_confirmed(profile, confirmed, primary_samples)


def default_mapping(
    registry: SchemeRegistry, scheme_id: str | None, preset_id: str | None, labels: Sequence[str]
) -> tuple[dict[str, str], tuple[str, ...]]:
    """(raw -> category, ordered targets) for a preset, or identity when no preset."""
    if preset_id:
        preset = registry.preset(preset_id)
        mapping = preset.resolved_rules(labels)
        return mapping, preset.targets
    identity = {label: label for label in labels}
    return identity, tuple(labels)


def presets_for(registry: SchemeRegistry, scheme_id: str | None) -> tuple[str, ...]:
    if not scheme_id:
        return ()
    return tuple(item.id for item in registry.presets_for(scheme_id))


def validate_comparison(spec: ComparisonSpec, labels_a: Sequence[str], labels_b: Sequence[str]):
    """Return human-readable problems (empty list when the comparison is complete)."""
    problems: list[str] = []
    if spec.kind == "reference":
        if not spec.categories:
            problems.append("공통 범주가 비어 있습니다.")
        for side, mapping, labels in (
            ("A", spec.mapping_a, labels_a),
            ("B", spec.mapping_b, labels_b),
        ):
            lookup = {k.casefold(): v for k, v in mapping.items()}
            missing = [
                label for label in labels if label.casefold() not in lookup and "*" not in mapping
            ]
            if missing:
                problems.append(f"{side}면 라벨에 매핑이 없습니다: {', '.join(missing)}")
            bad = [
                target
                for target in mapping.values()
                if target != EXCLUDE and target not in spec.categories
            ]
            if bad:
                problems.append(f"{side}면 매핑 대상이 공통 범주 밖입니다: {', '.join(bad)}")
    for cell in spec.cells:
        if not cell.labels_a or not cell.labels_b:
            problems.append(f"관심 셀 '{cell.name}'의 A/B 라벨을 선택하세요.")
    return problems


def evaluate(
    profile: ReportProfile,
    datasets: Mapping[str, ParsedDataset],
    registry: SchemeRegistry | None = None,
) -> ProfileResult:
    return profile_module.evaluate_profile(profile, datasets, registry=registry)


def generate(
    profile: ReportProfile,
    datasets: Mapping[str, ParsedDataset],
    output_path: str | Path,
    *,
    cancel_event: Event | None = None,
    progress: ProgressCallback | None = None,
    registry: SchemeRegistry | None = None,
) -> tuple[Path, Path | None]:
    """Evaluate, render, verify and commit the workbook(s); also saves the profile JSON."""
    from .workbook import generate_workbook, generate_workbook_pair, verify_workbook

    cancel = cancel_event or Event()
    if progress:
        progress(10, "비교 지표를 계산하는 중…")
    result = profile_module.evaluate_profile(profile, datasets, registry=registry)
    if result.blocked:
        raise RuntimeError("평가가 차단되었습니다: " + "; ".join(result.errors))
    if cancel.is_set():
        raise GenerationCancelled()
    if progress:
        progress(45, "Excel 통합문서를 작성하는 중…")
    output = Path(output_path)
    if profile.options.color_only_copy:
        with_codes, color_only = generate_workbook_pair(output, result, cancel_check=cancel)
    else:
        with_codes, color_only = generate_workbook(output, result, cancel_check=cancel), None
    if progress:
        progress(90, "통합문서를 다시 열어 검증하는 중…")
    verify_workbook(with_codes)
    if color_only is not None:
        verify_workbook(color_only)
    profile.save(output.with_suffix(".profile.json"))
    if progress:
        progress(100, "완료")
    return with_codes, color_only
