"""Workbook rendering for profile-driven (N-dataset, selectable-comparison) reports.

Sheet order::

    README → Label_Audit → Alignment → Joined_Data → S01…Snn (spatial maps, optional)
           → C01…Cnn (one per comparison) → Overall Summary

The legacy two-input renderer in :mod:`workbook` is untouched; this module is
selected when the evaluation object is a :class:`profile.ProfileResult`.
"""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from colorsys import hsv_to_rgb
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from . import association as assoc
from .core import _normalise
from .profile import (
    AssociationStats,
    ComparisonResult,
    ProfileResult,
    ReferenceStats,
    SampleData,
    _map_value,
)
from .workbook import (
    MAP_BORDER,
    WorkbookCancelled,
    _cancelled,
    _fill,
    _format_metric,
    _set_section,
    _style_header_row,
    _title,
    _write_legend,
)

REPORT_MODE = "profile"
MAP_TOP = 8  # first map heading row; each block is grid.rows + MAP_GAP rows tall
MAP_GAP = 10
KNOWN_STYLES: dict[str, tuple[str, str]] = {
    # legacy electrical
    "pass": ("P", "00FFFF"),
    "noactive": ("NA", "00B0F0"),
    "short": ("S", "FFC000"),
    "none": ("N", "FF0000"),
    "open": ("O", "FF00FF"),
    "nogateeffect": ("NG", "7030A0"),
    # electrical E-5
    "enormal": ("EN", "00FFFF"),
    "enoactive": ("ENA", "00B0F0"),
    "eopen": ("EO", "FF00FF"),
    "eshort": ("ES", "FFC000"),
    "einvalid": ("EI", "7030A0"),
    # ML / binary / optical
    "normal": ("N", "00FFFF"),
    "fail": ("F", "FF0000"),
    "good": ("G", "92D050"),
    "bad": ("B", "FF0000"),
    "exclude": ("X", "BFBFBF"),
    "missing": ("-", "F2F2F2"),
}
OUTCOME_STYLES = {
    "Match": ("=", "92D050"),
    "Mismatch": ("≠", "FF0000"),
    "Excluded": ("X", "BFBFBF"),
    "Missing": ("-", "F2F2F2"),
}
COOCCURRENCE_STYLES = {
    "Both": ("AB", "C00000"),
    "A only": ("A", "FFC000"),
    "B only": ("B", "00B0F0"),
    "Neither": ("·", "F2F2F2"),
    "Excluded": ("X", "BFBFBF"),
    "Missing": ("-", "FFFFFF"),
}
RESIDUAL_FILL_POS = "F8CBAD"
RESIDUAL_FILL_NEG = "BDD7EE"
REFERENCES = (
    "Agresti (2013) Categorical Data Analysis 3e — chi-square, odds ratio, relative risk, Fisher.",
    "Bergsma (2013) J. Korean Stat. Soc. 42:323–328 — bias-corrected Cramér's V.",
    "Haberman (1973) Biometrics 29:205–220 — adjusted standardized residuals.",
    "Theil (1970) Am. J. Sociol. 76:103–154 — uncertainty coefficient U.",
    "Cohen (1960) Educ. Psychol. Meas. 20:37–46 — kappa; Cicchetti & Feinstein (1990) "
    "J. Clin. Epidemiol. 43:551–558 — positive/negative agreement.",
    "Gorodkin (2004) Comput. Biol. Chem. 28:367–374 — multi-class MCC.",
    "Holm (1979) Scand. J. Stat. 6:65–70 — step-down multiple-test adjustment.",
)


# ------------------------------------------------------------------------ styling


def label_styles(labels: Sequence[str]) -> dict[str, tuple[str, str]]:
    """Code/colour per raw label: known vocabularies keep their palette, others get HSV."""
    styles: dict[str, tuple[str, str]] = {}
    used_codes: set[str] = set()
    for index, label in enumerate(labels):
        known = KNOWN_STYLES.get(_normalise(label))
        if known and known[0] not in used_codes:
            styles[label] = known
        else:
            rgb = hsv_to_rgb((index + 0.5) / max(len(labels), 1), 0.45, 0.90)
            color = "".join(f"{round(channel * 255):02X}" for channel in rgb)
            styles[label] = (f"U{index + 1}", color)
        used_codes.add(styles[label][0])
    return styles


def _style_map_block(ws: Any, heading_row: int, title: str, rows: int, nodes: int) -> int:
    """Axes/borders for one map whose body starts at ``heading_row + 1``; returns body row."""
    _set_section(ws.cell(heading_row - 1, 3), title)
    top_left = ws.cell(heading_row, 3, "Row / Node")
    top_left.font = Font(bold=True, size=8)
    top_left.border = MAP_BORDER
    top_left.alignment = Alignment(horizontal="center", vertical="center")
    body_row = heading_row + 1
    for node in range(1, nodes + 1):
        cell = ws.cell(heading_row, 3 + node, node)
        cell.font = Font(bold=True, size=8)
        cell.border = MAP_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(1, rows + 1):
        label = ws.cell(body_row + row_index - 1, 3, row_index)
        label.font = Font(bold=True, size=8)
        label.border = MAP_BORDER
        label.alignment = Alignment(horizontal="center", vertical="center")
        for node in range(1, nodes + 1):
            cell = ws.cell(body_row + row_index - 1, 3 + node)
            cell.border = MAP_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(3, 4 + nodes):
        ws.column_dimensions[get_column_letter(col)].width = 2.5
    for row in range(heading_row, body_row + rows):
        ws.row_dimensions[row].height = 17
    return body_row


def _paint(
    ws: Any,
    body_row: int,
    cells: Mapping[tuple[int, int], tuple[str, str]],
    *,
    show_codes: bool,
) -> None:
    for (row_index, node_index), (code, color) in cells.items():
        cell = ws.cell(body_row + row_index - 1, 3 + node_index)
        cell.value = code if show_codes else None
        cell.fill = _fill(color)
        cell.font = Font(size=7, bold=True)


def _counts_panel(
    ws: Any, row: int, col: int, title: str, styles: Mapping[str, tuple[str, str]], counts: Counter
) -> None:
    _write_legend(ws, row, col, title, styles)
    ws.cell(row, col + 2, "Count").font = Font(bold=True)
    ws.cell(row, col + 3, "Share").font = Font(bold=True)
    total = sum(counts.values())
    for offset, label in enumerate(styles, 1):
        ws.cell(row + offset, col + 2, counts.get(label, 0))
        share = ws.cell(row + offset, col + 3, counts.get(label, 0) / total if total else None)
        if share.value is not None:
            share.number_format = "0.0%"
        else:
            share.value = "N/A"


# ------------------------------------------------------------------------- helpers


def _fmt(cell: Any, value: Any, number_format: str = "0.000") -> None:
    if value is None:
        cell.value = "N/A"
    else:
        cell.value = float(value)
        cell.number_format = number_format


def _pvalue(cell: Any, value: float | None) -> None:
    if value is None:
        cell.value = "N/A"
    elif value == 0:
        cell.value = "< 1E-300"  # underflow: the exact tail is below double precision
    elif value < 1e-4:
        cell.value = value
        cell.number_format = "0.00E+00"
    else:
        cell.value = value
        cell.number_format = "0.0000"


def _outcome(result: ComparisonResult, raw_a: str | None, raw_b: str | None) -> str:
    if raw_a is None or raw_b is None:
        return "Missing"
    spec = result.spec
    cat_a, drop_a = _map_value(raw_a, spec.mapping_a, spec.exclude_a)
    cat_b, drop_b = _map_value(raw_b, spec.mapping_b, spec.exclude_b)
    if drop_a or drop_b:
        return "Excluded"
    if cat_a is None or cat_b is None:
        return "Missing"
    return "Match" if cat_a == cat_b else "Mismatch"


def _categories(result: ComparisonResult, raw_a: str | None, raw_b: str | None):
    spec = result.spec
    cat_a = cat_b = None
    if raw_a is not None:
        cat_a, drop = _map_value(raw_a, spec.mapping_a, spec.exclude_a)
        cat_a = "Exclude" if drop else cat_a
    if raw_b is not None:
        cat_b, drop = _map_value(raw_b, spec.mapping_b, spec.exclude_b)
        cat_b = "Exclude" if drop else cat_b
    return cat_a, cat_b


def _cooccurrence(result: ComparisonResult, cell_index: int, raw_a, raw_b) -> str:
    if raw_a is None or raw_b is None:
        return "Missing"
    spec = result.spec
    cat_a, cat_b = _categories(result, raw_a, raw_b)
    if cat_a == "Exclude" or cat_b == "Exclude":
        return "Excluded"
    if cat_a is None or cat_b is None:
        return "Missing"
    cell = spec.cells[cell_index]
    keys_a = {_normalise(x) for x in cell.labels_a}
    keys_b = {_normalise(x) for x in cell.labels_b}
    in_a = _normalise(cat_a) in keys_a
    in_b = _normalise(cat_b) in keys_b
    if in_a and in_b:
        return "Both"
    if in_a:
        return "A only"
    if in_b:
        return "B only"
    return "Neither"


def _dataset_title(result: ProfileResult, dataset_id: str) -> str:
    return result.profile.dataset(dataset_id).label


# -------------------------------------------------------------------------- sheets


def _readme(wb: Workbook, result: ProfileResult, *, show_map_codes: bool) -> None:
    ws = wb.active
    ws.title = "README"
    profile = result.profile
    _title(ws, "R2R Evaluation Workbook — profile report", "H")
    ws["A2"], ws["B2"] = "Report mode", REPORT_MODE
    ws["A3"], ws["B3"] = "Profile title", profile.options.title or "(untitled)"
    grid = result.datasets[profile.primary.id].grid
    ws["A4"], ws["B4"] = "Grid (rows x nodes)", f"{grid.rows}x{grid.nodes}"
    ws["A5"], ws["B5"] = "Samples aligned", len(result.samples)
    ws["A6"], ws["B6"] = "Blocked", "yes" if result.blocked else "no"
    ws["A7"], ws["B7"] = (
        "Map display mode",
        "Codes + fills" if show_map_codes else "Color fills only; map body codes blank",
    )
    for row in range(2, 8):
        ws.cell(row, 1).font = Font(bold=True)
    row = 9
    _set_section(ws.cell(row, 1), "Datasets")
    headers = ("Id", "Title", "Role", "Scheme", "Grid", "Path", "SHA256", "Worksheet", "Model")
    for col, header in enumerate(headers, 1):
        ws.cell(row + 1, col, header)
    _style_header_row(ws, row + 1, 1, len(headers))
    row += 2
    for spec in profile.datasets:
        for source in result.sources.get(spec.id, ()):
            values = (
                spec.id,
                spec.label,
                spec.role,
                spec.scheme or "(unregistered)",
                f"{spec.grid_rows}x{spec.grid_nodes}",
                source.path,
                source.sha256,
                source.worksheet or "CSV",
                spec.model_column or "",
            )
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            row += 1
    row += 1
    _set_section(ws.cell(row, 1), "Comparisons")
    headers = ("Id", "Kind", "Title", "Side A", "Side B", "Categories", "Preset A", "Preset B")
    for col, header in enumerate(headers, 1):
        ws.cell(row + 1, col, header)
    _style_header_row(ws, row + 1, 1, len(headers))
    row += 2
    for spec in profile.comparisons:
        values = (
            spec.id,
            spec.kind,
            spec.label,
            _dataset_title(result, spec.a),
            _dataset_title(result, spec.b),
            ", ".join(spec.categories) if spec.categories else "(raw labels)",
            spec.preset_a or "",
            spec.preset_b or "",
        )
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        row += 1
    row += 1
    _set_section(ws.cell(row, 1), "Errors / blocking conditions")
    if result.errors:
        for error in result.errors:
            row += 1
            ws.cell(row, 1, error).alignment = Alignment(wrap_text=True)
    else:
        row += 1
        ws.cell(row, 1, "none")
    row += 2
    _set_section(ws.cell(row, 1), "Semantics")
    notes = (
        "reference comparison: side A is ground truth for side B; confusion matrix rows are "
        "actual (A) and columns predicted (B). Majority baseline = share of A's largest class.",
        "association comparison: neither side is truth; symmetric measures only. Adjusted "
        "residual |z| > 1.96 marks a cell whose count departs from independence at ~5 %.",
        "Exclude: the record is dropped from that comparison and counted under 'excluded'. "
        "N/A: undefined denominator; never substituted by zero.",
        "Cochran warning: >20 % of expected counts < 5 or any < 1 — read chi-square with care; "
        "2x2 cells always carry Fisher's exact p.",
    )
    for note in notes:
        row += 1
        ws.cell(row, 1, note).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 30
    row += 2
    _set_section(ws.cell(row, 1), "References")
    for reference in REFERENCES:
        row += 1
        ws.cell(row, 1, reference)
    row += 2
    _set_section(ws.cell(row, 1), "Report profile (JSON; reproduce with --profile)")
    for line in profile.dumps().splitlines():
        row += 1
        ws.cell(row, 1, line).font = Font(name="Consolas", size=9)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28
    for col in "CDE":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["F"].width = 60
    ws.column_dimensions["G"].width = 40
    ws.column_dimensions["H"].width = 16
    ws.sheet_view.showGridLines = False


def _label_audit(wb: Workbook, result: ProfileResult) -> None:
    ws = wb.create_sheet("Label_Audit")
    _title(ws, "Label audit — raw label distribution per dataset", "H")
    headers = ("Dataset", "Scheme", "Raw label", "In scheme", "Count", "Share")
    for col, header in enumerate(headers, 1):
        ws.cell(3, col, header)
    _style_header_row(ws, 3, 1, len(headers))
    row = 4
    for spec in result.profile.datasets:
        dataset = result.datasets[spec.id]
        counts = Counter(
            dataset.canonical(record.value) for sheet in dataset.sheets for record in sheet.records
        )
        total = sum(counts.values())
        extra = {_normalise(x) for x in result.unregistered.get(spec.id, ())}
        for label in result.labels(spec.id):
            ws.cell(row, 1, spec.label)
            ws.cell(row, 2, spec.scheme or "(unregistered)")
            ws.cell(row, 3, label)
            ws.cell(row, 4, "no" if _normalise(label) in extra else "yes")
            ws.cell(row, 5, counts[label])
            ws.cell(row, 6, counts[label] / total if total else None).number_format = "0.00%"
            row += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{max(row - 1, 4)}"


def _alignment(wb: Workbook, result: ProfileResult) -> None:
    ws = wb.create_sheet("Alignment")
    _title(ws, "Sample alignment — which sheet of each dataset forms one sample", "H")
    ids = [spec.id for spec in result.profile.datasets]
    headers = ("Sample", *[_dataset_title(result, dataset_id) for dataset_id in ids])
    for col, header in enumerate(headers, 1):
        ws.cell(3, col, header)
    _style_header_row(ws, 3, 1, len(headers))
    for row, sample in enumerate(result.samples, 4):
        ws.cell(row, 1, sample.sample)
        for col, dataset_id in enumerate(ids, 2):
            ws.cell(row, col, sample.members.get(dataset_id, "(not aligned)"))
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 30
    ws.freeze_panes = "B4"


def _joined(wb: Workbook, result: ProfileResult) -> None:
    ws = wb.create_sheet("Joined_Data")
    _title(ws, "Joined data — one row per aligned coordinate", "L")
    ids = [spec.id for spec in result.profile.datasets]
    headers = ["Sample", "Row", "Node"]
    headers += [f"{_dataset_title(result, dataset_id)} raw" for dataset_id in ids]
    for comparison in result.comparisons:
        prefix = comparison.spec.id
        headers += [f"{prefix}: A", f"{prefix}: B", f"{prefix}: outcome"]
        if comparison.spec.kind == "association":
            headers += [f"{prefix}: {cell.name}" for cell in comparison.spec.cells]
    for col, header in enumerate(headers, 1):
        ws.cell(3, col, header)
    _style_header_row(ws, 3, 1, len(headers))
    row = 4
    for sample in result.samples:
        for record in sample.records:
            values: list[Any] = [sample.sample, record.row, record.node]
            values += [record.value(dataset_id) for dataset_id in ids]
            for comparison in result.comparisons:
                raw_a = record.value(comparison.spec.a)
                raw_b = record.value(comparison.spec.b)
                cat_a, cat_b = _categories(comparison, raw_a, raw_b)
                outcome = (
                    _outcome(comparison, raw_a, raw_b)
                    if comparison.spec.kind == "reference"
                    else None
                )
                values += [cat_a, cat_b, outcome]
                if comparison.spec.kind == "association":
                    values += [
                        _cooccurrence(comparison, index, raw_a, raw_b)
                        for index in range(len(comparison.spec.cells))
                    ]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, "N/A" if value is None else value)
            row += 1
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(row - 1, 4)}"
    ws.column_dimensions["A"].width = 28
    for col in range(4, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18


def _sample_sheet(
    wb: Workbook,
    result: ProfileResult,
    sample: SampleData,
    index: int,
    styles: Mapping[str, Mapping[str, tuple[str, str]]],
    *,
    show_map_codes: bool,
) -> None:
    ws = wb.create_sheet(f"S{index:02d}")
    grid = sample.grid
    _title(ws, f"R2R Spatial Maps — {sample.sample}", get_column_letter(4 + grid.nodes + 12))
    ws["A2"] = "Sample"
    ws["B2"] = sample.sample
    panel_col = 4 + grid.nodes + 2
    heading = MAP_TOP
    block = grid.rows + MAP_GAP
    breaks: list[int] = []
    coords = {(record.row, record.node): record for record in sample.records}

    def next_block() -> int:
        nonlocal heading
        current = heading
        heading += block
        breaks.append(heading - 3)
        return current

    for spec in result.profile.datasets:
        if spec.id not in sample.members:
            continue
        top = next_block()
        body = _style_map_block(ws, top, f"{spec.label} — raw labels", grid.rows, grid.nodes)
        palette = styles[spec.id]
        cells = {}
        counts: Counter = Counter()
        for key, record in coords.items():
            value = record.value(spec.id)
            if value is None:
                continue
            cells[key] = palette[value]
            counts[value] += 1
        _paint(ws, body, cells, show_codes=show_map_codes)
        _counts_panel(ws, top, panel_col, f"{spec.label} legend", palette, counts)
    for comparison in result.comparisons:
        spec = comparison.spec
        if spec.a not in sample.members or spec.b not in sample.members:
            continue
        if spec.kind == "reference":
            top = next_block()
            body = _style_map_block(
                ws, top, f"{spec.label} — agreement (A vs B)", grid.rows, grid.nodes
            )
            cells = {}
            counts = Counter()
            for key, record in coords.items():
                outcome = _outcome(comparison, record.value(spec.a), record.value(spec.b))
                cells[key] = OUTCOME_STYLES[outcome]
                counts[outcome] += 1
            _paint(ws, body, cells, show_codes=show_map_codes)
            _counts_panel(ws, top, panel_col, "Agreement legend", OUTCOME_STYLES, counts)
        else:
            for cell_index, cell_spec in enumerate(spec.cells):
                top = next_block()
                body = _style_map_block(
                    ws,
                    top,
                    f"{spec.label} — co-occurrence: {cell_spec.name}",
                    grid.rows,
                    grid.nodes,
                )
                cells = {}
                counts = Counter()
                for key, record in coords.items():
                    state = _cooccurrence(
                        comparison, cell_index, record.value(spec.a), record.value(spec.b)
                    )
                    cells[key] = COOCCURRENCE_STYLES[state]
                    counts[state] += 1
                _paint(ws, body, cells, show_codes=show_map_codes)
                _counts_panel(
                    ws, top, panel_col, "Co-occurrence legend", COOCCURRENCE_STYLES, counts
                )
    ws.column_dimensions[get_column_letter(panel_col)].width = 6
    ws.column_dimensions[get_column_letter(panel_col + 1)].width = 34
    ws.column_dimensions[get_column_letter(panel_col + 2)].width = 9
    ws.column_dimensions[get_column_letter(panel_col + 3)].width = 9
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    last_row = heading - MAP_GAP + 2
    ws.print_area = f"A1:{get_column_letter(panel_col + 4)}{last_row}"
    for break_row in breaks[:-1]:
        ws.row_breaks.append(Break(id=break_row))


def _write_table(ws: Any, row: int, col: int, title: str, headers, rows) -> int:
    _set_section(ws.cell(row, col), title)
    for offset, header in enumerate(headers):
        ws.cell(row + 1, col + offset, header)
    _style_header_row(ws, row + 1, col, col + len(headers) - 1)
    for r, values in enumerate(rows, row + 2):
        for offset, value in enumerate(values):
            cell = ws.cell(r, col + offset)
            if isinstance(value, tuple) and len(value) == 2 and isinstance(value[1], str):
                _fmt(cell, value[0], value[1])
            elif value is None:
                cell.value = "N/A"
            else:
                cell.value = value
    return row + 2 + len(rows)


def _write_crosstab(
    ws: Any,
    row: int,
    col: int,
    title: str,
    table: assoc.ContingencyTable,
    *,
    corner: str,
    values: Sequence[Sequence[Any]] | None = None,
    number_format: str | None = None,
    highlight: Callable[[Any], str | None] | None = None,
    totals: bool = True,
) -> int:
    _set_section(ws.cell(row, col), title)
    ws.cell(row + 1, col, corner)
    for j, label in enumerate(table.labels_b, 1):
        ws.cell(row + 1, col + j, label)
    if totals:
        ws.cell(row + 1, col + len(table.labels_b) + 1, "Total")
    _style_header_row(ws, row + 1, col, col + len(table.labels_b) + (1 if totals else 0))
    data = values if values is not None else table.counts
    for i, label in enumerate(table.labels_a):
        r = row + 2 + i
        ws.cell(r, col, label).font = Font(bold=True)
        for j in range(len(table.labels_b)):
            cell = ws.cell(r, col + 1 + j)
            value = data[i][j]
            if value is None:
                cell.value = "N/A"
            else:
                cell.value = value
                if number_format:
                    cell.number_format = number_format
            if highlight is not None and (color := highlight(value)):
                cell.fill = _fill(color)
            cell.border = MAP_BORDER
            cell.alignment = Alignment(horizontal="center")
        if totals:
            ws.cell(r, col + len(table.labels_b) + 1, table.row_totals[i]).font = Font(bold=True)
    end = row + 2 + len(table.labels_a)
    if totals:
        ws.cell(end, col, "Total").font = Font(bold=True)
        for j, total in enumerate(table.column_totals, 1):
            ws.cell(end, col + j, total).font = Font(bold=True)
        ws.cell(end, col + len(table.labels_b) + 1, table.total).font = Font(bold=True)
        end += 1
    return end


def _reference_block(ws: Any, row: int, title: str, stats: ReferenceStats) -> int:
    metrics = stats.metrics
    table = assoc.ContingencyTable(metrics.labels, metrics.labels, metrics.matrix)
    end = _write_crosstab(ws, row, 1, f"{title} — confusion matrix", table, corner="Actual \\ Pred")
    agreement = stats.agreement
    general = [
        ("Pairs scored", stats.pairs),
        ("Excluded records", stats.excluded),
        ("Accuracy", (metrics.accuracy, "0.0%")),
        (
            "Accuracy CI95",
            None
            if metrics.accuracy_ci95 is None
            else f"{metrics.accuracy_ci95[0]:.1%} – {metrics.accuracy_ci95[1]:.1%}",
        ),
        ("Majority-class baseline", (agreement.majority_baseline, "0.0%")),
        ("Balanced accuracy", (metrics.balanced_accuracy, "0.0%")),
        ("Macro F1", (metrics.macro_f1, "0.000")),
        ("Weighted F1", (metrics.weighted_f1, "0.000")),
        ("Cohen's kappa", (agreement.kappa, "0.000")),
        ("MCC", (agreement.mcc, "0.000")),
        ("Positive agreement", (agreement.positive_agreement, "0.000")),
        ("Negative agreement", (agreement.negative_agreement, "0.000")),
    ]
    end = _write_table(ws, end + 1, 1, f"{title} — summary metrics", ("Metric", "Value"), general)
    per_class = [
        (
            item.label,
            item.support,
            (item.precision, "0.000"),
            (item.recall, "0.000"),
            (item.f1, "0.000"),
            (item.specificity, "0.000"),
            (item.false_positive_rate, "0.000"),
            (item.false_negative_rate, "0.000"),
        )
        for item in metrics.per_class
    ]
    end = _write_table(
        ws,
        end + 1,
        1,
        f"{title} — per class (one-vs-rest)",
        ("Class", "Support", "Precision", "Recall", "F1", "Specificity", "FPR", "FNR"),
        per_class,
    )
    end = _write_crosstab(
        ws, end + 1, 1, f"{title} — raw label cross-table", stats.raw_table, corner="A raw \\ B raw"
    )
    return end


def _association_block(ws: Any, row: int, title: str, stats: AssociationStats) -> int:
    table = stats.table
    end = _write_crosstab(ws, row, 1, f"{title} — counts", table, corner="A \\ B")
    row_pct = [
        [None if rt == 0 else c / rt for c in r]
        for r, rt in zip(table.counts, table.row_totals, strict=True)
    ]
    end = _write_crosstab(
        ws,
        end + 1,
        1,
        f"{title} — row % (share of each A label)",
        table,
        corner="A \\ B",
        values=row_pct,
        number_format="0.0%",
        totals=False,
    )
    col_pct = [
        [
            None if table.column_totals[j] == 0 else c / table.column_totals[j]
            for j, c in enumerate(r)
        ]
        for r in table.counts
    ]
    end = _write_crosstab(
        ws,
        end + 1,
        1,
        f"{title} — column % (share of each B label)",
        table,
        corner="A \\ B",
        values=col_pct,
        number_format="0.0%",
        totals=False,
    )
    if stats.chi.expected:
        end = _write_crosstab(
            ws,
            end + 1,
            1,
            f"{title} — expected counts under independence",
            table,
            corner="A \\ B",
            values=stats.chi.expected,
            number_format="0.0",
            totals=False,
        )
    end = _write_crosstab(
        ws,
        end + 1,
        1,
        f"{title} — adjusted standardized residuals (Haberman)",
        table,
        corner="A \\ B",
        values=stats.residuals,
        number_format="0.00",
        totals=False,
        highlight=lambda z: None
        if z is None or abs(z) <= 1.96
        else (RESIDUAL_FILL_POS if z > 0 else RESIDUAL_FILL_NEG),
    )
    chi = stats.chi
    theil = stats.theil
    summary = [
        ("Pairs", stats.pairs),
        ("Excluded records", stats.excluded),
        ("Chi-square", (chi.statistic, "0.00")),
        ("Degrees of freedom", chi.dof),
        ("Chi-square p", chi.p_value),
        ("Cramér's V", (chi.cramers_v, "0.000")),
        ("Cramér's V (Bergsma bias-corrected)", (chi.cramers_v_corrected, "0.000")),
        ("Min expected count", (chi.min_expected, "0.00")),
        ("Share of expected < 5", (chi.low_expected_fraction, "0.0%")),
        ("Cochran warning", "yes" if chi.cochran_warning else "no"),
        ("Theil's U(B|A) — A explains B", (theil.u_b_given_a, "0.000")),
        ("Theil's U(A|B) — B explains A", (theil.u_a_given_b, "0.000")),
    ]
    if stats.agreement is not None:
        summary += [
            ("Cohen's kappa (shared categories)", (stats.agreement.kappa, "0.000")),
            ("Observed agreement", (stats.agreement.observed_agreement, "0.0%")),
            ("Positive agreement", (stats.agreement.positive_agreement, "0.000")),
            ("Negative agreement", (stats.agreement.negative_agreement, "0.000")),
        ]
    end = _write_table(
        ws, end + 1, 1, f"{title} — association summary", ("Metric", "Value"), summary
    )
    for offset, (label, _value) in enumerate(summary):
        if label == "Chi-square p":
            _pvalue(ws.cell(end - len(summary) + offset, 2), chi.p_value)
        if label == "Degrees of freedom":
            ws.cell(end - len(summary) + offset, 2).number_format = "0"
    if stats.cells:
        rows = []
        for cell in stats.cells:
            rows.append(
                (
                    cell.name,
                    cell.a,
                    cell.b,
                    cell.c,
                    cell.d,
                    (cell.odds_ratio, "0.000"),
                    None
                    if cell.odds_ratio_ci95 is None
                    else f"{cell.odds_ratio_ci95[0]:.3f} – {cell.odds_ratio_ci95[1]:.3f}",
                    (cell.relative_risk, "0.000"),
                    (cell.phi, "0.000"),
                    (cell.yule_q, "0.000"),
                    cell.fisher_p,
                    cell.holm_p,
                    "yes" if cell.corrected else "no",
                )
            )
        start = end + 1
        end = _write_table(
            ws,
            start,
            1,
            f"{title} — 2x2 cells of interest (a=both, b=A only, c=B only, d=neither)",
            (
                "Cell",
                "a",
                "b",
                "c",
                "d",
                "Odds ratio",
                "OR CI95 (Woolf)",
                "Relative risk",
                "phi",
                "Yule's Q",
                "Fisher p",
                "Holm p",
                "0.5 corrected",
            ),
            rows,
        )
        for r in range(start + 2, end):
            _pvalue(ws.cell(r, 11), ws.cell(r, 11).value if ws.cell(r, 11).value != "N/A" else None)
            _pvalue(ws.cell(r, 12), ws.cell(r, 12).value if ws.cell(r, 12).value != "N/A" else None)
    return end


def _comparison_sheet(wb: Workbook, result: ProfileResult, comparison: ComparisonResult, index):
    ws = wb.create_sheet(f"C{index:02d}")
    spec = comparison.spec
    _title(ws, f"R2R Comparison — {spec.label} [{spec.kind}]", "P")
    ws["A2"], ws["B2"] = "Comparison id", spec.id
    ws["A3"], ws["B3"] = "Side A", _dataset_title(result, spec.a)
    ws["A4"], ws["B4"] = "Side B", _dataset_title(result, spec.b)
    ws["A5"], ws["B5"] = (
        "Reading",
        "A is ground truth; B is evaluated against it."
        if spec.kind == "reference"
        else "Symmetric association; neither side is truth.",
    )
    for row in range(2, 6):
        ws.cell(row, 1).font = Font(bold=True)
    row = 7
    if comparison.overall is None:
        _set_section(ws.cell(row, 1), "Not evaluated")
        ws.cell(row + 1, 1, "Unmapped side A labels: " + (", ".join(comparison.unmapped_a) or "—"))
        ws.cell(row + 2, 1, "Unmapped side B labels: " + (", ".join(comparison.unmapped_b) or "—"))
        return
    mapping_rows = [(raw, target, "A") for raw, target in spec.mapping_a.items()]
    mapping_rows += [(raw, target, "B") for raw, target in spec.mapping_b.items()]
    mapping_rows += [(raw, "Exclude", "A") for raw in spec.exclude_a]
    mapping_rows += [(raw, "Exclude", "B") for raw in spec.exclude_b]
    if mapping_rows:
        row = (
            _write_table(
                ws, row, 1, "Label mapping applied", ("Raw label", "Category", "Side"), mapping_rows
            )
            + 1
        )
    block = _reference_block if spec.kind == "reference" else _association_block
    row = block(ws, row, "Overall", comparison.overall) + 1
    if len(comparison.per_sample) > 1:
        ws.row_breaks.append(Break(id=row - 1))
        if spec.kind == "reference":
            rows = [
                (
                    name,
                    stats.pairs,
                    stats.excluded,
                    (stats.metrics.accuracy, "0.0%"),
                    (stats.metrics.macro_f1, "0.000"),
                    (stats.agreement.kappa, "0.000"),
                    (stats.agreement.mcc, "0.000"),
                    (stats.agreement.majority_baseline, "0.0%"),
                )
                for name, stats in comparison.per_sample.items()
            ]
            row = _write_table(
                ws,
                row,
                1,
                "Per sample",
                ("Sample", "Pairs", "Excluded", "Accuracy", "Macro F1", "Kappa", "MCC", "Baseline"),
                rows,
            )
        else:
            headers = ["Sample", "Pairs", "Excluded", "Chi-square p", "Cramér's V (bc)"]
            headers += [f"OR: {cell.name}" for cell in spec.cells]
            rows = []
            for name, stats in comparison.per_sample.items():
                values: list[Any] = [
                    name,
                    stats.pairs,
                    stats.excluded,
                    stats.chi.p_value,
                    (stats.chi.cramers_v_corrected, "0.000"),
                ]
                values += [(cell.odds_ratio, "0.000") for cell in stats.cells]
                rows.append(tuple(values))
            start = row
            row = _write_table(ws, row, 1, "Per sample", headers, rows)
            for r in range(start + 2, row):
                value = ws.cell(r, 4).value
                _pvalue(ws.cell(r, 4), None if value == "N/A" else value)
        row += 1
        for name, stats in comparison.per_sample.items():
            ws.row_breaks.append(Break(id=row - 1))
            row = block(ws, row, f"Sample {name}", stats) + 1
    ws.column_dimensions["A"].width = 38
    for col in range(2, 16):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:P{max(row, 20)}"


def _summary(wb: Workbook, result: ProfileResult) -> None:
    ws = wb.create_sheet("Overall Summary")
    _title(ws, "R2R Overall Summary", "L")
    ws["A2"] = (
        f"Datasets: {len(result.profile.datasets)} | Samples: {len(result.samples)} | "
        f"Comparisons: {len(result.comparisons)}"
    )
    headers = (
        "Comparison",
        "Kind",
        "Pairs",
        "Excluded",
        "Accuracy / Cramér's V (bc)",
        "Baseline / chi-square p",
        "Macro F1 / Theil U(B|A)",
        "Kappa",
        "MCC / Theil U(A|B)",
    )
    for col, header in enumerate(headers, 1):
        ws.cell(3, col, header)
    _style_header_row(ws, 3, 1, len(headers))
    row = 4
    for comparison in result.comparisons:
        stats = comparison.overall
        ws.cell(row, 1, comparison.spec.label)
        ws.cell(row, 2, comparison.spec.kind)
        if stats is None:
            ws.cell(row, 3, "not evaluated")
        elif isinstance(stats, ReferenceStats):
            ws.cell(row, 3, stats.pairs)
            ws.cell(row, 4, stats.excluded)
            _format_metric(ws.cell(row, 5), stats.metrics.accuracy)
            _format_metric(ws.cell(row, 6), stats.agreement.majority_baseline)
            _fmt(ws.cell(row, 7), stats.metrics.macro_f1)
            _fmt(ws.cell(row, 8), stats.agreement.kappa)
            _fmt(ws.cell(row, 9), stats.agreement.mcc)
        else:
            ws.cell(row, 3, stats.pairs)
            ws.cell(row, 4, stats.excluded)
            _fmt(ws.cell(row, 5), stats.chi.cramers_v_corrected)
            _pvalue(ws.cell(row, 6), stats.chi.p_value)
            _fmt(ws.cell(row, 7), stats.theil.u_b_given_a)
            _fmt(ws.cell(row, 8), None if stats.agreement is None else stats.agreement.kappa)
            _fmt(ws.cell(row, 9), stats.theil.u_a_given_b)
        row += 1
    row += 1
    _set_section(ws.cell(row, 1), "Label distribution per dataset")
    dist_header = row + 1
    for col, header in enumerate(("Dataset", "Label", "Count", "Share"), 1):
        ws.cell(dist_header, col, header)
    _style_header_row(ws, dist_header, 1, 4)
    row = dist_header + 1
    first_dataset_rows: tuple[int, int] | None = None
    for spec in result.profile.datasets:
        dataset = result.datasets[spec.id]
        counts = Counter(
            dataset.canonical(record.value) for sheet in dataset.sheets for record in sheet.records
        )
        total = sum(counts.values())
        start = row
        for label in result.labels(spec.id):
            ws.cell(row, 1, spec.label)
            ws.cell(row, 2, label)
            ws.cell(row, 3, counts[label])
            ws.cell(row, 4, counts[label] / total if total else None).number_format = "0.00%"
            row += 1
        if first_dataset_rows is None and row > start:
            first_dataset_rows = (start, row - 1)
    if first_dataset_rows is not None:
        start, end = first_dataset_rows
        chart = BarChart()
        chart.title = f"{result.profile.primary.label} label counts"
        chart.title.overlay = False
        chart.legend = None
        chart.y_axis.scaling.min = 0
        chart.y_axis.delete = chart.x_axis.delete = False
        chart.add_data(Reference(ws, min_col=3, min_row=start, max_row=end))
        chart.set_categories(Reference(ws, min_col=2, min_row=start, max_row=end))
        chart.width, chart.height = 16, 9
        ws.add_chart(chart, "K4")
    ws.column_dimensions["A"].width = 40
    for col in range(2, 10):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:T{max(row, 30)}"


# ------------------------------------------------------------------------- driver


def render_profile_workbook(
    wb: Workbook,
    result: ProfileResult,
    cancel_check: Callable[[], bool] | Any | None = None,
    *,
    show_map_codes: bool = True,
) -> None:
    if result.blocked:
        raise ValueError("Cannot render a blocked profile report: " + "; ".join(result.errors))
    _readme(wb, result, show_map_codes=show_map_codes)
    _label_audit(wb, result)
    _alignment(wb, result)
    if result.profile.options.joined_data:
        _joined(wb, result)
    if result.profile.options.spatial_maps:
        styles = {spec.id: label_styles(result.labels(spec.id)) for spec in result.profile.datasets}
        for index, sample in enumerate(result.samples, 1):
            if _cancelled(cancel_check):
                raise WorkbookCancelled("Workbook generation was cancelled")
            _sample_sheet(wb, result, sample, index, styles, show_map_codes=show_map_codes)
    for index, comparison in enumerate(result.comparisons, 1):
        if _cancelled(cancel_check):
            raise WorkbookCancelled("Workbook generation was cancelled")
        _comparison_sheet(wb, result, comparison, index)
    _summary(wb, result)


def verify_profile_workbook(wb: Workbook) -> None:
    """Structural verification of a profile report after reopening it."""
    names = wb.sheetnames
    readme = wb["README"]
    if readme["B2"].value != REPORT_MODE:
        raise ValueError("README!B2 must record the profile report mode")
    expected_head = ["README", "Label_Audit", "Alignment"]
    if names[: len(expected_head)] != expected_head:
        raise ValueError("README, Label_Audit, Alignment must be first")
    if names[-1] != "Overall Summary":
        raise ValueError("Overall Summary must be last")
    middle = names[len(expected_head) : -1]
    if middle and middle[0] == "Joined_Data":
        joined = wb["Joined_Data"]
        if joined.freeze_panes != "D4" or joined.auto_filter.ref is None:
            raise ValueError("Joined_Data must retain freeze panes and filter")
        middle = middle[1:]
    samples = [name for name in middle if name.startswith("S")]
    comparisons = [name for name in middle if name.startswith("C")]
    if samples + comparisons != middle:
        raise ValueError("Sample map sheets must precede comparison sheets")
    if samples != [f"S{i:02d}" for i in range(1, len(samples) + 1)]:
        raise ValueError("Sample sheets must be ordered Snn identifiers")
    if comparisons != [f"C{i:02d}" for i in range(1, len(comparisons) + 1)]:
        raise ValueError("Comparison sheets must be ordered Cnn identifiers")
    rows_text, nodes_text = str(readme["B4"].value).split("x")
    rows, nodes = int(rows_text), int(nodes_text)
    show_codes = readme["B7"].value == "Codes + fills"
    for name in samples:
        ws = wb[name]
        if ws.cell(MAP_TOP, 3).value != "Row / Node":
            raise ValueError(f"{name}: first map header missing")
        if ws.cell(MAP_TOP, 3 + nodes).value != nodes or ws.cell(MAP_TOP + rows, 3).value != rows:
            raise ValueError(f"{name}: incomplete {rows}x{nodes} axes")
        for row in ws.iter_rows(
            min_row=MAP_TOP + 1, max_row=MAP_TOP + rows, min_col=4, max_col=3 + nodes
        ):
            for cell in row:
                if cell.fill.patternType != "solid" or (cell.value is not None) != show_codes:
                    raise ValueError(f"{name}: invalid map cell {cell.coordinate}")
        if ws.page_setup.orientation != "landscape":
            raise ValueError(f"{name}: expected landscape")
    for name in comparisons:
        ws = wb[name]
        if not str(ws["A1"].value).startswith("R2R Comparison — "):
            raise ValueError(f"{name}: comparison title missing")
        if ws["B2"].value is None:
            raise ValueError(f"{name}: comparison id missing")
    summary = wb["Overall Summary"]
    if summary["A3"].value != "Comparison":
        raise ValueError("Overall Summary header missing")


def profile_json_from_readme(wb: Workbook) -> dict[str, Any]:
    """Recover the embedded profile JSON from a rendered README sheet."""
    ws = wb["README"]
    lines: list[str] = []
    collecting = False
    for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
        value = row[0]
        if collecting:
            if value is None:
                continue
            lines.append(str(value))
        elif isinstance(value, str) and value.startswith("Report profile (JSON"):
            collecting = True
    return json.loads("\n".join(lines))
