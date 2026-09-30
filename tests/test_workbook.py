from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import r2r_evaluation_report.workbook as workbook_module
from r2r_evaluation_report.core import (
    ParsedDataset,
    ParsedSheet,
    SourceMetadata,
    build_synthetic_evaluation,
    evaluate,
)
from r2r_evaluation_report.workbook import (
    AGREEMENT_STYLES,
    JOINED_HEADERS,
    MEASUREMENT_STYLES,
    PREDICTION_STYLES,
    WorkbookCancelled,
    derive_color_only_path,
    generate_workbook,
    generate_workbook_pair,
    verify_workbook,
)


def _records_from_synthetic():
    base = build_synthetic_evaluation()
    joined = base.sheet_evaluations[0].joined
    measurements = tuple(item.measurement for item in joined if item.measurement)
    predictions = tuple(item.prediction for item in joined if item.prediction)
    return base, measurements, predictions


def _two_sample_evaluation():
    base, measurements, predictions = _records_from_synthetic()
    digest = base.measurement_sources[0].sha256
    titles = (
        "Synthetic SAM 11",
        "Very long printed-device sample name retained without tab truncation",
    )
    measurement_sheets = []
    prediction_sheets = []
    for index, title in enumerate(titles, 1):
        measurement_source = SourceMetadata(
            f"<synthetic-measurement-{index}>",
            digest if index == 1 else "c" * 64,
            index * 100,
            index * 1_000_000_000,
            "synthetic",
            title,
        )
        prediction_source = SourceMetadata(
            f"<synthetic-prediction-{index}>",
            digest if index == 1 else "d" * 64,
            index * 200,
            index * 2_000_000_000,
            "synthetic",
            title,
        )
        measurement_sheets.append(ParsedSheet("measurement", measurement_source, measurements, {}))
        prediction_sheets.append(ParsedSheet("prediction", prediction_source, predictions, {}))
    return evaluate(
        ParsedDataset("measurement", "<measurements>", tuple(measurement_sheets)),
        ParsedDataset("prediction", "<predictions>", tuple(prediction_sheets)),
    )


@pytest.fixture(scope="module")
def evaluation():
    return _two_sample_evaluation()


@pytest.fixture(scope="module")
def generated_report(tmp_path_factory, evaluation):
    output = tmp_path_factory.mktemp("workbook") / "evaluation.xlsx"
    generate_workbook(output, evaluation)
    return output


def _series_formula(series, kind: str) -> str:
    source = series.val if kind == "value" else series.cat
    for name in ("numRef", "strRef"):
        reference = getattr(source, name, None)
        if reference is not None and reference.f:
            return reference.f
    return ""


def _find_value(ws, value, min_row=1, max_row=None, min_col=1, max_col=None):
    max_row = max_row or ws.max_row
    max_col = max_col or ws.max_column
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            if cell.value == value:
                return cell
    raise AssertionError(f"{value!r} was not found in {ws.title}")


def test_core_result_sources_profile_order_and_joined_contract(
    generated_report,
    evaluation,
):
    verify_workbook(generated_report)
    wb = load_workbook(generated_report)
    assert wb.sheetnames == [
        "README",
        "Mapping_Audit",
        "Joined_Data",
        "R01",
        "R02",
        "Overall Summary",
    ]

    readme = wb["README"]
    for source in (*evaluation.measurement_sources, *evaluation.prediction_sources):
        path_cell = _find_value(readme, source.path)
        row = path_cell.row
        assert readme.cell(row, 3).value == source.sha256
        assert readme.cell(row, 4).value == source.size
        assert readme.cell(row, 5).value == source.mtime_ns
        assert readme.cell(row, 6).value == source.encoding
        assert readme.cell(row, 7).value == source.worksheet
    no_gate = _find_value(readme, "nogateeffect")
    assert readme.cell(no_gate.row, 2).value == "Exclude"
    assert readme.cell(no_gate.row, 3).value == "Fail"
    assert "missing pair" in readme["B4"].value
    profile_header = _find_value(readme, "Expanded Normal scenario")
    assert profile_header.fill.fgColor.rgb[-6:] == "1F4E78"

    audit = wb["Mapping_Audit"]
    long_name = evaluation.sheet_evaluations[1].measurement_sheet
    assert _find_value(audit, long_name).value == long_name
    assert _find_value(audit, "exact").value == "exact"

    joined = wb["Joined_Data"]
    assert (
        tuple(joined.cell(3, col).value for col in range(1, len(JOINED_HEADERS) + 1))
        == JOINED_HEADERS
    )
    assert joined.max_column == 16
    assert joined.max_row == 3 + 2 * 26 * 38
    assert joined.freeze_panes == "A4"
    assert joined.auto_filter.ref == f"A3:P{joined.max_row}"
    assert joined["D9"].value == "No Gate Effect"
    assert joined["E9"].value == "Exclude"
    assert joined["J9"].value == "Fail"
    assert joined["K9"].value == "Pass"
    assert joined["L9"].value == "FN"
    assert joined["M9"].value == 7
    assert joined["N9"].value == "Normal"
    wb.close()


def test_report_all_grid_cells_colors_page_bands_metrics_and_charts(
    generated_report,
    evaluation,
):
    wb = load_workbook(generated_report)
    report = wb["R01"]
    sample = evaluation.sheet_evaluations[0]
    assert sample.measurement_sheet in report["A1"].value
    assert report["D8"].value == 1 and report["AO8"].value == 38
    assert report["C9"].value == 1 and report["C34"].value == 26
    assert report["D43"].value == 1 and report["AO43"].value == 38
    assert report["C44"].value == 1 and report["C69"].value == 26
    assert report["D78"].value == 1 and report["AO78"].value == 38
    assert report["C79"].value == 1 and report["C104"].value == 26

    for item in sample.joined:
        measurement = item.measurement
        prediction = item.prediction
        assert measurement is not None and prediction is not None
        row, node = measurement.row, measurement.node
        measurement_cell = report.cell(8 + row, 3 + node)
        prediction_cell = report.cell(43 + row, 3 + node)
        agreement_cell = report.cell(78 + row, 3 + node)
        measurement_code, measurement_color = MEASUREMENT_STYLES[measurement.value]
        prediction_code, prediction_color = PREDICTION_STYLES[prediction.value]
        actual_binary = evaluation.status_mapping.binary[
            "".join(character for character in measurement.value.casefold() if character.isalnum())
        ]
        predicted_binary = "Pass" if prediction.value == "Normal" else "Fail"
        agreement = {
            ("Fail", "Fail"): "TP",
            ("Pass", "Pass"): "TN",
            ("Pass", "Fail"): "FP",
            ("Fail", "Pass"): "FN",
        }[(actual_binary, predicted_binary)]
        agreement_code, agreement_color = AGREEMENT_STYLES[agreement]
        assert measurement_cell.value == measurement_code
        assert measurement_cell.fill.fgColor.rgb[-6:] == measurement_color
        assert prediction_cell.value == prediction_code
        assert prediction_cell.fill.fgColor.rgb[-6:] == prediction_color
        assert agreement_cell.value == agreement_code
        assert agreement_cell.fill.fgColor.rgb[-6:] == agreement_color

    assert report["I79"].value == "FN"
    assert report["I79"].fill.fgColor.rgb[-6:] == "FF0000"
    assert report["AZ4"].value == "Measurement legend"
    assert report["AZ43"].value == "Prediction legend"
    assert report["BH76"].value == "Agreement legend"
    assert report["AW113"].value == "Class"
    assert report["BH113"].value == "Expanded Normal legend"
    assert report["AQ101"].value.startswith("Definitions: Fail is the positive")
    assert report.max_row == 142 and report.max_column == 65

    assert report.column_dimensions["C"].width == 2.5
    assert report.column_dimensions["AO"].width == 2.5
    for heading, body in ((8, 9), (43, 44), (78, 79), (113, 114)):
        assert all(report.row_dimensions[row].height == 17 for row in range(heading, body + 26))
    assert report.print_area == "'R01'!$A$1:$BM$142"
    assert report.page_setup.orientation == "landscape"
    assert str(report.page_setup.paperSize) == str(report.PAPERSIZE_A3)
    assert report.page_setup.fitToWidth == 1 and report.page_setup.fitToHeight == 0
    assert [item.id for item in report.row_breaks.brk] == [39, 74, 109]

    raw = sample.raw_status_by_prediction
    assert report["AQ76"].value == "Raw status x prediction"
    assert [report.cell(77, col).value for col in range(44, 47)] == [
        "Normal",
        "Open",
        "Short",
    ]
    no_gate_row = _find_value(report, "No Gate Effect", 78, 85, 43, 43).row
    assert [report.cell(no_gate_row, col).value for col in range(44, 47)] == [
        raw["No Gate Effect"].get(label, 0) for label in ("Normal", "Open", "Short")
    ]
    assert report["AW76"].value == "Mapped 3x3"
    assert report["BC76"].value == "Binary 2x2"
    assert report["AR87"].value == pytest.approx(sample.operational_binary.accuracy)
    assert report["AR92"].value == pytest.approx(sample.yield_stats.measurement_pass_rate)
    fail = next(item for item in sample.operational_binary.per_class if item.label == "Fail")
    assert report["AS99"].value == pytest.approx(fail.recall)
    assert report["AT99"].value == pytest.approx(fail.f1)
    assert report["AV99"].value == pytest.approx(fail.specificity)
    assert report["AW99"].value == pytest.approx(fail.false_positive_rate)
    assert report["AX99"].value == pytest.approx(fail.false_negative_rate)
    assert report["AS93"].value == pytest.approx(sample.yield_stats.measurement_pass_ci95[0])
    assert report["AT93"].value == pytest.approx(sample.yield_stats.measurement_pass_ci95[1])

    assert len(report._charts) == 2
    comparison, distribution = report._charts
    assert [_series_formula(series, "value") for series in comparison.series] == [
        "'R01'!$AR$5:$AR$6",
        "'R01'!$AS$5:$AS$6",
    ]
    assert _series_formula(comparison.series[0], "category") == "'R01'!$AQ$5:$AQ$6"
    assert [_series_formula(series, "value") for series in distribution.series] == [
        "'R01'!$AR$44:$AR$48",
        "'R01'!$AS$44:$AS$48",
    ]
    assert comparison.anchor._from.row == 10
    assert distribution.anchor._from.row == 50
    wb.close()


def test_overall_summary_uses_authoritative_metrics_and_has_no_overlap(
    generated_report,
    evaluation,
):
    wb = load_workbook(generated_report)
    summary = wb["Overall Summary"]
    assert summary["A4"].value == "OVERALL"
    assert summary["B4"].value == evaluation.overall_binary.total
    assert summary["C4"].value == evaluation.overall_three_class.total
    assert summary["F4"].value == pytest.approx(evaluation.overall_binary.accuracy)
    fail = next(item for item in evaluation.overall_binary.per_class if item.label == "Fail")
    assert summary["M4"].value == pytest.approx(fail.recall)
    assert summary["N4"].value == pytest.approx(fail.f1)
    assert summary["O4"].value == pytest.approx(evaluation.overall_yield.measurement_pass_rate)
    assert summary["P4"].value == pytest.approx(evaluation.overall_yield.prediction_pass_rate)
    section_row = 4 + len(evaluation.sheet_evaluations) + 22
    assert summary.cell(section_row, 1).value == "Overall mapped 3x3"
    assert summary.cell(section_row, 7).value == "Overall binary 2x2"
    assert summary.cell(section_row, 13).value == "Overall raw status x prediction"
    rank_title = _find_value(summary, "Fail recall / F1 rank")
    ranked_names = {
        summary.cell(row, rank_title.column + 1).value
        for row in range(rank_title.row + 2, rank_title.row + 2 + len(evaluation.sheet_evaluations))
    }
    assert ranked_names == {sample.measurement_sheet for sample in evaluation.sheet_evaluations}
    assert _find_value(summary, "Exclusions / unmatched audit").row == rank_title.row
    assert len(summary._charts) == 1
    chart = summary._charts[0]
    assert [_series_formula(series, "value") for series in chart.series] == [
        "'Overall Summary'!$O$4:$O$6",
        "'Overall Summary'!$P$4:$P$6",
    ]
    assert summary.page_setup.orientation == "landscape"
    assert str(summary.page_setup.paperSize) == str(summary.PAPERSIZE_A3)
    assert summary.page_setup.fitToWidth == 1 and summary.page_setup.fitToHeight == 0
    assert str(summary.print_area).startswith("'Overall Summary'!$A$1:$Q$")
    scenario_title = _find_value(summary, "Expanded Normal scenario — per-sample KPI")
    assert [item.id for item in summary.row_breaks.brk] == [section_row - 1, scenario_title.row - 1]
    assert scenario_title.column == 1
    assert _find_value(summary, "Overall Expanded Normal scenario 3x3").column == 1
    wb.close()


def test_custom_status_profile_drives_joined_canonical_and_binary(tmp_path):
    base, measurements, predictions = _records_from_synthetic()
    custom_measurements = (replace(measurements[0], value="Custom Ink Gap"), *measurements[1:])
    title = "Custom mapping sample"
    measured_source = replace(
        base.measurement_sources[0], path="<custom-measurement>", worksheet=title
    )
    predicted_source = replace(
        base.prediction_sources[0], path="<custom-prediction>", worksheet=title
    )
    evaluation = evaluate(
        ParsedDataset(
            "measurement",
            measured_source.path,
            (ParsedSheet("measurement", measured_source, custom_measurements, {}),),
        ),
        ParsedDataset(
            "prediction",
            predicted_source.path,
            (ParsedSheet("prediction", predicted_source, predictions, {}),),
        ),
        status_rules={"Custom Ink Gap": "Open"},
        binary_status_rules={"Custom Ink Gap": "Fail"},
        expanded_normal_status_rules={"Custom Ink Gap": "Open"},
    )
    assert not evaluation.blocked
    output = generate_workbook(tmp_path / "custom.xlsx", evaluation)
    wb = load_workbook(output)
    joined = wb["Joined_Data"]
    assert joined["D4"].value == "Custom Ink Gap"
    assert joined["E4"].value == "Open"
    assert joined["J4"].value == "Fail"
    profile = wb["README"]
    custom = _find_value(profile, "custominkgap")
    assert profile.cell(custom.row, 2).value == "Open"
    assert profile.cell(custom.row, 3).value == "Fail"
    wb.close()


def test_summary_layout_scales_past_ten_samples_without_overlap():
    base = build_synthetic_evaluation()
    first = base.sheet_evaluations[0]
    samples = tuple(
        replace(first, measurement_sheet=f"Scaled sample {index:02d}") for index in range(1, 13)
    )
    evaluation = replace(base, sheet_evaluations=samples)
    wb = Workbook()
    wb.remove(wb.active)
    workbook_module._summary_sheet(wb, evaluation, samples)
    ws = wb["Overall Summary"]
    assert ws["A16"].value == "Scaled sample 12"
    assert ws["A38"].value == "Overall mapped 3x3"
    rank_title = _find_value(ws, "Fail recall / F1 rank")
    names = [
        ws.cell(row, rank_title.column + 1).value
        for row in range(rank_title.row + 2, rank_title.row + 14)
    ]
    assert set(names) == {f"Scaled sample {index:02d}" for index in range(1, 13)}
    assert all(name is not None for name in names)
    wb.close()


class _CancelAfterVerification:
    def __init__(self):
        self.calls = 0

    def __call__(self):
        self.calls += 1
        return self.calls >= 4


def test_atomic_overwrite_cancel_error_cleanup_and_reopen(tmp_path, monkeypatch):
    evaluation = build_synthetic_evaluation()
    output = generate_workbook(tmp_path / "atomic.xlsx", evaluation)
    original = output.read_bytes()

    cancellation = _CancelAfterVerification()
    with pytest.raises(WorkbookCancelled):
        generate_workbook(output, evaluation, cancel_check=cancellation)
    assert output.read_bytes() == original
    assert not (tmp_path / ".atomic.partial.xlsx").exists()

    def fail_verification(_path):
        raise ValueError("forced structural verification failure")

    monkeypatch.setattr(workbook_module, "verify_workbook_structure", fail_verification)
    with pytest.raises(ValueError, match="forced structural"):
        generate_workbook(output, evaluation)
    assert output.read_bytes() == original
    assert not (tmp_path / ".atomic.partial.xlsx").exists()
    reopened = load_workbook(output)
    assert reopened.sheetnames[-1] == "Overall Summary"
    reopened.close()


def test_pair_generation_preserves_map_fills_and_blanks_only_body_codes(tmp_path):
    evaluation = build_synthetic_evaluation()
    output = tmp_path / "긴 한글 결과" / "평가.xlsx"
    output.parent.mkdir()
    text_path, color_path = generate_workbook_pair(output, evaluation)
    assert text_path == output
    assert color_path == derive_color_only_path(output)
    text_wb = load_workbook(text_path)
    color_wb = load_workbook(color_path)
    try:
        assert text_wb.sheetnames == color_wb.sheetnames
        for name in text_wb.sheetnames:
            if name.startswith("R") and name[1:].isdigit():
                text_ws, color_ws = text_wb[name], color_wb[name]
                body_ranges = {
                    (row, col)
                    for body in (9, 44, 79, 114)
                    for row in range(body, body + 26)
                    for col in range(4, 42)
                }
                for row in range(1, max(text_ws.max_row, color_ws.max_row) + 1):
                    for col in range(1, max(text_ws.max_column, color_ws.max_column) + 1):
                        if (row, col) not in body_ranges:
                            assert color_ws.cell(row, col).value == text_ws.cell(row, col).value
                        assert repr(color_ws.cell(row, col).fill) == repr(
                            text_ws.cell(row, col).fill
                        )
                        assert repr(color_ws.cell(row, col).border) == repr(
                            text_ws.cell(row, col).border
                        )
                for body in (9, 44, 79, 114):
                    for row in range(body, body + 26):
                        for col in range(4, 42):
                            assert color_ws.cell(row, col).value is None
                            assert (
                                color_ws.cell(row, col).fill.fgColor.rgb
                                == text_ws.cell(row, col).fill.fgColor.rgb
                            )
                assert color_ws["A1"].value == text_ws["A1"].value
                assert color_ws["D4"].value == text_ws["D4"].value
                assert color_ws["AZ4"].value == text_ws["AZ4"].value == "Measurement legend"
                assert color_ws["AZ43"].value == text_ws["AZ43"].value == "Prediction legend"
                assert color_ws["BH76"].value == text_ws["BH76"].value == "Agreement legend"
                assert color_ws["BH113"].value == text_ws["BH113"].value == "Expanded Normal legend"
        assert text_wb["README"]["B7"].value == "Codes + fills"
        assert "Color fills only" in color_wb["README"]["B7"].value
    finally:
        text_wb.close()
        color_wb.close()


def test_pair_cancel_keeps_both_existing_outputs_and_cleans_temps(tmp_path):
    evaluation = build_synthetic_evaluation()
    output = tmp_path / "atomic.xlsx"
    color = derive_color_only_path(output)
    output.write_bytes(b"old text")
    color.write_bytes(b"old color")
    with pytest.raises(WorkbookCancelled):
        generate_workbook_pair(output, evaluation, cancel_check=lambda: True)
    assert output.read_bytes() == b"old text"
    assert color.read_bytes() == b"old color"
    assert not list(tmp_path.glob(".*partial.xlsx"))
    assert not list(tmp_path.glob(".*backup.xlsx"))


def test_derive_color_only_path_normalizes_xlsx_and_adds_missing_suffix():
    assert derive_color_only_path(Path("C:/긴 경로/평가.XLSX")) == Path(
        "C:/긴 경로/평가-color-only.xlsx"
    )
    assert derive_color_only_path(Path("C:/긴 경로/평가")) == Path(
        "C:/긴 경로/평가-color-only.xlsx"
    )


def test_pair_stale_backup_guard_preserves_recoverable_backup(tmp_path):
    evaluation = build_synthetic_evaluation()
    output = tmp_path / "atomic.xlsx"
    backup = tmp_path / ".atomic.backup.xlsx"
    output.write_bytes(b"existing")
    backup.write_bytes(b"recoverable")
    with pytest.raises(RuntimeError, match="Stale transaction backup"):
        generate_workbook_pair(output, evaluation)
    assert output.read_bytes() == b"existing"
    assert backup.read_bytes() == b"recoverable"
    assert not list(tmp_path.glob(".*partial.xlsx"))


@pytest.mark.parametrize("color_preexisting", [False, True])
def test_pair_second_commit_failure_rolls_back_existing_and_missing_outputs(
    tmp_path, monkeypatch, color_preexisting
):
    evaluation = build_synthetic_evaluation()
    output = tmp_path / "atomic.xlsx"
    color = derive_color_only_path(output)
    output.write_bytes(b"old text")
    if color_preexisting:
        color.write_bytes(b"old color")
    original_replace = workbook_module.os.replace

    def fail_color_commit(source, destination):
        if Path(destination) == color and ".partial" in Path(source).name:
            raise OSError("forced second commit failure")
        return original_replace(source, destination)

    monkeypatch.setattr(workbook_module.os, "replace", fail_color_commit)
    with pytest.raises(OSError, match="second commit"):
        generate_workbook_pair(output, evaluation)
    assert output.read_bytes() == b"old text"
    assert color.read_bytes() == b"old color" if color_preexisting else not color.exists()
    assert not list(tmp_path.glob(".*partial.xlsx"))
    assert not list(tmp_path.glob(".*backup.xlsx"))
