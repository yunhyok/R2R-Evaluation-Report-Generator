"""Rendering and re-open verification of profile-driven workbooks."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from r2r_evaluation_report import profile, schemes
from r2r_evaluation_report.profile_workbook import (
    label_styles,
    profile_json_from_readme,
    verify_profile_workbook,
)
from r2r_evaluation_report.workbook import generate_workbook, generate_workbook_pair

E5 = ("E-Normal", "E-NoActive", "E-Open", "E-Short", "E-Invalid")
GRID = [(r, n) for r in range(1, 27) for n in range(1, 39)]


def _write(path: Path, name: str, column: str, values) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["Name", "Row", "Node", column])
        writer.writeheader()
        for index, (row, node) in enumerate(GRID):
            writer.writerow({"Name": name, "Row": row, "Node": node, column: values(index)})
    return path


@pytest.fixture(scope="module")
def result(tmp_path_factory) -> profile.ProfileResult:
    root = tmp_path_factory.mktemp("profile-workbook")

    def elec(index):
        return E5[index % 5]

    def optical(index):
        if elec(index) in {"E-Invalid", "E-Short"}:
            return "BAD" if index % 3 else "GOOD"
        return "OPEN" if elec(index) == "E-Open" and index % 2 else "GOOD"

    def ml(index):
        base = {"E-Normal": "Normal", "E-NoActive": "Normal", "E-Open": "Open"}
        return base.get(elec(index), "Short") if index % 13 else "Normal"

    registry = schemes.load_registry(user_path="/nonexistent")
    e5 = registry.scheme("electrical_e5").labels
    datasets = (
        profile.DatasetSpec(
            "elec",
            str(_write(root / "e.csv", "20260731_7kgf_sam1", "Status", elec)),
            "electrical_gt",
            "electrical_e5",
            title="Electrical",
        ),
        profile.DatasetSpec(
            "human",
            str(_write(root / "h.csv", "20260731_7kgf_sam1", "label", optical)),
            "optical_human",
            "optical_3",
            title="Human optical",
        ),
        profile.DatasetSpec(
            "ml",
            str(_write(root / "m.csv", "20260731_7kgf_sam1", "prediction", ml)),
            "electrical_ml",
            "ml_3class",
            title="ML",
        ),
    )
    comparisons = (
        profile.ComparisonSpec(
            "self",
            "reference",
            "elec",
            "ml",
            "Electrical vs ML",
            mapping_a=registry.preset("e5_to_3class").resolved_rules(e5),
            mapping_b={label: label for label in ("Normal", "Open", "Short")},
            categories=("Normal", "Open", "Short"),
            preset_a="e5_to_3class",
        ),
        profile.ComparisonSpec(
            "cross",
            "association",
            "elec",
            "human",
            "Electrical vs optical",
            cells=(profile.CellSpec("E-Invalid x BAD", ("E-Invalid",), ("BAD",)),),
        ),
    )
    spec = profile.ReportProfile(datasets, comparisons, options=profile.ProfileOptions(title="t"))
    return profile.evaluate_profile(spec, registry=registry)


def test_profile_workbook_pair_renders_verifies_and_embeds_profile(tmp_path, result) -> None:
    assert not result.blocked, result.errors
    with_codes, color_only = generate_workbook_pair(tmp_path / "report.xlsx", result)
    wb = load_workbook(with_codes)
    assert wb.sheetnames == [
        "README",
        "Label_Audit",
        "Alignment",
        "Joined_Data",
        "S01",
        "C01",
        "C02",
        "Overall Summary",
    ]
    verify_profile_workbook(wb)
    embedded = profile.ReportProfile.from_dict(profile_json_from_readme(wb))
    assert embedded.comparisons == result.profile.comparisons
    assert embedded.datasets == result.profile.datasets
    maps = wb["S01"]
    # 3 raw-label maps + 1 agreement map + 1 co-occurrence map, each 26 + 10 rows tall.
    headings = [
        maps.cell(row, 3).value
        for row in range(1, maps.max_row + 1)
        if isinstance(maps.cell(row, 3).value, str) and maps.cell(row, 3).value != "Row / Node"
    ]
    assert len(headings) == 5 and headings[-1].endswith("E-Invalid x BAD")
    assert maps["C8"].value == "Row / Node" and maps["AO8"].value == 38 and maps["C34"].value == 26
    assert maps["D9"].value is not None
    fills_only = load_workbook(color_only)
    verify_profile_workbook(fills_only)
    assert fills_only["S01"]["D9"].value is None
    assert fills_only["S01"]["D9"].fill.fgColor.rgb == maps["D9"].fill.fgColor.rgb
    reference = wb["C01"]
    values = {
        reference.cell(row, 1).value: reference.cell(row, 2).value
        for row in range(1, reference.max_row + 1)
    }
    assert values["Pairs scored"] == 988 - 197 and values["Excluded records"] == 197
    assert values["Majority-class baseline"] == pytest.approx(396 / 791)
    association = wb["C02"]
    labels = [association.cell(row, 1).value for row in range(1, association.max_row + 1)]
    assert "Overall — adjusted standardized residuals (Haberman)" in labels
    assert "Overall — 2x2 cells of interest (a=both, b=A only, c=B only, d=neither)" in labels
    summary = wb["Overall Summary"]
    assert summary["A4"].value == "Electrical vs ML" and summary["B5"].value == "association"
    assert len(summary._charts) == 1


def test_profile_workbook_honours_options_and_blocks(tmp_path, result) -> None:
    trimmed = profile.ReportProfile(
        result.profile.datasets,
        (),
        result.profile.alignments,
        profile.ProfileOptions(spatial_maps=False, joined_data=False),
    )
    evaluated = profile.evaluate_profile(trimmed, result.datasets)
    output = generate_workbook(tmp_path / "trimmed.xlsx", evaluated)
    wb = load_workbook(output)
    assert wb.sheetnames == ["README", "Label_Audit", "Alignment", "Overall Summary"]
    verify_profile_workbook(wb)
    blocked = profile.evaluate_profile(
        profile.ReportProfile(
            (profile.DatasetSpec("elec", result.profile.datasets[0].path, scheme="ml_3class"),)
        )
    )
    assert blocked.blocked
    with pytest.raises(ValueError, match="blocked"):
        generate_workbook(tmp_path / "blocked.xlsx", blocked)
    assert not (tmp_path / "blocked.xlsx").exists()


def test_label_styles_keep_known_palettes_and_unique_codes() -> None:
    styles = label_styles(("E-Normal", "E-Invalid", "Custom A", "Custom B"))
    assert styles["E-Normal"] == ("EN", "00FFFF") and styles["E-Invalid"] == ("EI", "7030A0")
    assert styles["Custom A"][0] != styles["Custom B"][0]
    assert len({code for code, _ in styles.values()}) == 4
