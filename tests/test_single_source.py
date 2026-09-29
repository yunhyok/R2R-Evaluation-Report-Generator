from __future__ import annotations

import csv
from dataclasses import replace
from threading import Event

import pytest
from openpyxl import Workbook, load_workbook

from r2r_evaluation_report.core import DataContractError, describe_dataset, parse_dataset
from r2r_evaluation_report.gui import CoreWorkflowAdapter, MainWindow, WorkflowContext
from r2r_evaluation_report.workbook import generate_workbook_pair, verify_workbook
from scripts.create_synthetic_inputs import create_inputs


@pytest.mark.parametrize("kind", ["measurement", "prediction"])
@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_single_source_preserves_labels_without_comparison(tmp_path, kind, file_format):
    paths = dict(zip(("measurement", "prediction"), create_inputs(tmp_path), strict=True))
    source = paths[kind]
    with source.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    if kind == "measurement":
        # TXT Converter raw headers are lower-case and its rules allow custom labels.
        rows[0] = [header.lower() for header in rows[0]]
        rows[1][3] = "=Custom status"
    if file_format == "xlsx":
        source = source.with_suffix(".xlsx")
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.active["D2"].data_type = "s"
        wb.save(source)
        wb.close()
    else:
        with source.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(rows)
    parsed = parse_dataset(source, kind)
    result = describe_dataset(parsed)
    assert not result.blocked and result.report_mode == kind
    assert result.overall_three_class is result.overall_binary is result.overall_yield is None
    assert result.overall_expanded_normal is None
    assert all(sample.mapped_three_class is None for sample in result.sheet_evaluations)
    assert len(result.sheet_evaluations) == (2 if kind == "measurement" else 3)
    # Prediction-only samples are included when there is no measurement input.
    outputs = generate_workbook_pair(tmp_path / "report.xlsx", result)
    for index, output in enumerate(outputs):
        verify_workbook(output)
        wb = load_workbook(output)
        try:
            assert wb.sheetnames[:2] == ["README", "Source_Data"]
            assert "Mapping_Audit" not in wb and "Joined_Data" not in wb
            assert wb["Source_Data"].max_row - 1 == len(rows) - 1
            assert wb["Source_Data"]["D2"].value == rows[1][3]
            assert wb["Source_Data"]["D2"].data_type == "s"
            sheet = wb["R01"]
            assert sheet["D9"].value is None if index else sheet["D9"].value is not None
            assert sheet["D9"].fill.patternType == "solid"
            assert sheet["C43"].value is None
            assert not sheet._charts and not sheet.row_breaks.brk
            assert (
                sum(
                    sheet.cell(r, 45).value
                    for r in range(9, 9 + len(set(row[3] for row in rows[1:])))
                )
                == 988
            )
            summary = wb["Overall Summary"]
            assert len(summary._charts) == 1
            assert summary._charts[0].y_axis.scaling.min == 0
            assert (
                sum(
                    row[2]
                    for row in summary.iter_rows(min_row=4, max_col=4, values_only=True)
                    if row[0] == "OVERALL"
                )
                == len(rows) - 1
            )
            for worksheet in (sheet, summary, wb["Source_Data"]):
                text = str(list(worksheet.values)).casefold()
                for absent in ("macro-f1", "accuracy", "confusion", "agreement", "expanded normal"):
                    assert absent not in text
        finally:
            wb.close()


def test_single_source_keeps_coordinate_and_prediction_validation(tmp_path):
    measurement, prediction = create_inputs(tmp_path)
    prediction.write_text(
        prediction.read_text(encoding="utf-8-sig").replace(",Normal,", ",Other,"),
        encoding="utf-8-sig",
    )
    result = describe_dataset(parse_dataset(prediction, "prediction"))
    assert result.blocked and result.unknown_predictions == ("Other",)
    with pytest.raises(ValueError, match="blocked"):
        generate_workbook_pair(tmp_path / "blocked.xlsx", result)
    lines = measurement.read_text(encoding="utf-8-sig").splitlines()
    measurement.write_text("\n".join(lines[:-1]), encoding="utf-8-sig")
    with pytest.raises(DataContractError, match="988"):
        parse_dataset(measurement, "measurement")


@pytest.mark.parametrize("kind", ["measurement", "prediction"])
def test_single_input_gui_and_stale_preflight(qtbot, tmp_path, kind):
    measurement, prediction = create_inputs(tmp_path)
    window = MainWindow(test_mode=True)
    qtbot.addWidget(window)
    window.show()
    getattr(window, f"{kind}_edit").setText(
        str(measurement if kind == "measurement" else prediction)
    )
    window.output_edit.setText(str(tmp_path / "ui.xlsx"))
    assert not window.generate_button.isEnabled()
    window.start_preflight()
    qtbot.waitUntil(lambda: window._thread is None, timeout=10000)
    assert window.generate_button.isEnabled()
    assert not window.mapping_box.isVisible() and not window.label_box.isVisible()
    assert "단독" in window.summary_label.text()
    window.start_generation()
    qtbot.waitUntil(lambda: window._thread is None, timeout=30000)
    assert window.open_color_only_button.isVisible()
    verify_workbook(tmp_path / "ui.xlsx")
    other = "prediction" if kind == "measurement" else "measurement"
    getattr(window, f"{other}_edit").setText(
        str(prediction if other == "prediction" else measurement)
    )
    assert not window.generate_button.isEnabled()
    assert window.mapping_box.isVisible() and window.label_box.isVisible()
    window.start_preflight()
    qtbot.waitUntil(lambda: window._thread is None, timeout=10000)
    window._confirm_all_mappings()
    assert window.generate_button.isEnabled()
    getattr(window, f"{kind}_edit").clear()
    assert not window.generate_button.isEnabled()


def test_single_xlsx_selection_requires_new_preflight(tmp_path):
    measurement, _ = create_inputs(tmp_path)
    with measurement.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    source = tmp_path / "multiple.xlsx"
    wb = Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    for ws in wb:
        for row in rows:
            ws.append(row)
    wb.save(source)
    wb.close()
    adapter = CoreWorkflowAdapter()
    context = WorkflowContext(
        str(source), "", str(tmp_path / "out.xlsx"), None, None, [], [], Event(), lambda *_: None
    )
    assert not adapter.preflight(context).get("single_source_ready")
    with pytest.raises(RuntimeError, match="사전 검사"):
        adapter.generate(context)
    selected = replace(context, measurement_sheet="Second")
    assert adapter.preflight(selected)["single_source_ready"]
    with pytest.raises(RuntimeError, match="바뀌었습니다"):
        adapter.generate(replace(selected, measurement_sheet="First"))


def test_missing_input_and_input_overwrite_are_blocked(qtbot, tmp_path):
    window = MainWindow(test_mode=True)
    qtbot.addWidget(window)
    window.output_edit.setText(str(tmp_path / "out.xlsx"))
    assert "하나 이상" in window._validate_paths()
    source = tmp_path / "out-color-only.xlsx"
    source.touch()
    window.measurement_edit.setText(str(source))
    assert "다른 경로" in window._validate_paths()
