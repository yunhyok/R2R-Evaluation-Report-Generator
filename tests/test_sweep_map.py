"""Measured sweeps: parsing, the transfer-curve grid image, and its place in the report."""

from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from r2r_evaluation_report import core, profile, schemes
from r2r_evaluation_report.profile_workbook import (
    sweep_map_count,
    sweep_sources,
    verify_profile_workbook,
)
from r2r_evaluation_report.sweep_map import (
    CURRENT_FLOOR,
    EMPTY_FILL,
    log_current_bounds,
    render_sweep_grid,
)
from r2r_evaluation_report.workbook import generate_workbook_pair

GRID = [(r, n) for r in range(1, 27) for n in range(1, 39)]
VOLTAGES = tuple(float(v) for v in range(-20, 21, 2))
E5 = ("E-Normal", "E-NoActive", "E-Open", "E-Short", "E-Invalid")


def _current(label: str, voltage: float) -> float:
    if label == "E-Normal":
        return 5e-9 + 1e-6 / (1.0 + math.exp((voltage + 5.0) / 2.0))
    if label == "E-NoActive":
        return 3e-8
    if label == "E-Open":
        return 2e-12
    if label == "E-Short":
        return 1e-5
    return -3e-7 if voltage > 0 else 1e-7  # E-Invalid: bipolar swing


def _write_measurement(
    path: Path,
    names: tuple[str, ...],
    column: str = "Status",
    *,
    raw_layout: bool = False,
    blank: tuple[tuple[int, int], ...] = (),
) -> Path:
    headers = ["Name", "Row", "Node", column]
    headers += [f"I_{v:g}V" if raw_layout else f"{v:.5f}" for v in VOLTAGES]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for name in names:
            for index, (row, node) in enumerate(GRID):
                label = E5[index % 5]
                currents = [
                    "" if (row, node) in blank else f"{_current(label, v):.6e}" for v in VOLTAGES
                ]
                writer.writerow([name, row, node, label, *currents])
    return path


def test_sweep_columns_accept_merged_and_raw_layouts_only_as_monotonic_sets() -> None:
    merged = ["Name", "Row", "Node", "ON", "OFF", "ON/OFF", "gm", "Vth", "Status"]
    merged += ["-20.00000", "-19.80000", "0.000000", "20.00000"]
    assert core.sweep_columns(merged) == (
        ("-20.00000", -20.0),
        ("-19.80000", -19.8),
        ("0.000000", 0.0),
        ("20.00000", 20.0),
    )
    raw = ["name", "row", "node", "status", "classification_rule", "cond1", "I_-20V", "I_15V"]
    assert core.sweep_columns(raw) == (("I_-20V", -20.0), ("I_15V", 15.0))
    assert core.sweep_columns(["Name", "Row", "Node", "Status", -20.0, 15.0]) == (
        ("-20.0", -20.0),
        ("15.0", 15.0),
    )
    assert core.sweep_columns(["Name", "Row", "Node", "Status", "1"]) == ()
    assert core.sweep_columns(["Name", "Row", "Node", "Status", "-20", "15", "0"]) == ()
    assert core.sweep_columns(["Name", "Row", "Node", "label"]) == ()


def test_parse_dataset_keeps_sweeps_per_worksheet_with_nan_for_unreadable_cells(
    tmp_path: Path,
) -> None:
    source = _write_measurement(
        tmp_path / "merged.csv", ("sam1", "sam2"), blank=((1, 2),), raw_layout=False
    )
    dataset = core.parse_dataset(source, "label")
    first, second = dataset.sheets
    assert first.sweeps is not None and first.sweeps is second.sweeps
    table = first.sweeps
    assert table.voltages == VOLTAGES and table.points == 21
    assert table.currents.shape == (2 * 988, 21) and table.currents.dtype == np.float32
    assert first.records[0].sweep_index == 0 and second.records[0].sweep_index == 988
    trace = table.trace(first.records[0].sweep_index)
    assert trace is not None and trace[0] == pytest.approx(_current("E-Normal", -20.0), rel=1e-6)
    blank_record = next(r for r in first.records if r.coordinate == (1, 2))
    assert np.isnan(table.trace(blank_record.sweep_index)).all()
    assert table.trace(None) is None and table.trace(10_000) is None
    assert table.same_data(second.sweeps) and not table.same_data(
        core.SweepTable(VOLTAGES[:-1], table.currents[:, :-1], table.headers[:-1])
    )
    # The raw layout parses to the same axis; a file without sweep columns has none.
    raw = core.parse_dataset(
        _write_measurement(tmp_path / "raw.csv", ("sam1",), raw_layout=True), "label"
    )
    assert raw.sheets[0].sweeps is not None and raw.sheets[0].sweeps.voltages == VOLTAGES
    plain = tmp_path / "plain.csv"
    with plain.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Name", "Row", "Node", "label"])
        writer.writerows(["s", r, n, "GOOD"] for r, n in GRID)
    assert core.parse_dataset(plain, "label").sheets[0].sweeps is None
    assert core.parse_dataset(plain, "label").sheets[0].records[0].sweep_index is None
    # Sample inspection never reads sweeps; the xlsx reader does (numeric header cells).
    assert len(core.inspect_samples(source, "label")) == 2
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "150px"
    sheet.append(["Name", "Row", "Node", "Status", *VOLTAGES])
    for index, (row, node) in enumerate(GRID):
        label = E5[index % 5]
        sheet.append(["sam1", row, node, label, *[_current(label, v) for v in VOLTAGES]])
    workbook.save(tmp_path / "data.xlsx")
    parsed = core.parse_dataset(tmp_path / "data.xlsx", "label")
    assert parsed.sheets[0].sweeps is not None
    assert parsed.sheets[0].sweeps.voltages == VOLTAGES
    assert parsed.sheets[0].sweeps.currents[0, 0] == pytest.approx(
        _current("E-Normal", -20.0), rel=1e-6
    )


def test_render_sweep_grid_is_deterministic_and_marks_missing_cells() -> None:
    traces = {}
    for index, key in enumerate(GRID):
        if key == (2, 3):
            continue  # no sweep at all
        label = E5[index % 5]
        values = np.array([_current(label, v) for v in VOLTAGES], dtype=np.float32)
        if key == (2, 4):
            values[:] = np.nan  # unreadable sweep
        if key == (2, 5):
            values[3:6] = np.nan  # a gap inside the trace still draws the rest
        traces[key] = (values, {"E-Open": "FF00FF"}.get(label, "00FFFF"))
    first = render_sweep_grid(26, 38, VOLTAGES, traces)
    second = render_sweep_grid(26, 38, VOLTAGES, traces)
    assert first.png_bytes() == second.png_bytes()
    assert (first.width, first.height) == (26 + 38 * 32 + 4, 14 + 26 * 24 + 14)
    assert first.cells == 988 and first.traces == 988 - 2
    # E-Short traces sit at 1e-5 A, so the shared axis tops out at that decade.
    assert first.log_current_range == (-13, -5) and first.voltage_range == (-20.0, 20.0)
    assert "1e-13 to 1e-5 A" in first.axis_note

    def centre(row: int, node: int) -> tuple[int, int, int]:
        x = 26 + (node - 1) * 32 + 16
        y = 14 + (row - 1) * 24 + 12
        return first.image.getpixel((x, y))[:3]

    assert centre(2, 3) == EMPTY_FILL and centre(2, 4) == EMPTY_FILL
    assert centre(2, 5) != EMPTY_FILL
    # A present cell's curve pixels carry the label colour somewhere in the cell.
    x0, y0 = 26 + (1 - 1) * 32, 14 + (1 - 1) * 24
    colours = {
        first.image.getpixel((x, y))[:3] for x in range(x0, x0 + 32) for y in range(y0, y0 + 24)
    }
    assert any(abs(r - 0) < 80 and g > 150 and b > 150 for r, g, b in colours)

    assert log_current_bounds([np.array([np.nan, np.nan])]) == (-13, -12)
    assert log_current_bounds([np.array([CURRENT_FLOOR / 10, 2e-5])]) == (-13, -4)
    with pytest.raises(ValueError, match="two voltage"):
        render_sweep_grid(26, 38, (0.0,), {})
    with pytest.raises(ValueError, match="span"):
        render_sweep_grid(26, 38, (1.0, 1.0), {})


@pytest.fixture
def sweep_profile(tmp_path: Path) -> profile.ReportProfile:
    electrical = _write_measurement(tmp_path / "e.csv", ("20260731_7kgf_sam1",))
    # An ImageMarker-style write-back: the measurement workbook with a human label column
    # and the very same sweep columns -> must not draw the curves twice.
    writeback = _write_measurement(tmp_path / "h.csv", ("20260731_7kgf_sam1",), column="label")
    registry = schemes.load_registry(user_path="/nonexistent")
    optical = tmp_path / "o.csv"
    with optical.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Name", "Row", "Node", "label"])
        for index, (row, node) in enumerate(GRID):
            writer.writerow(["20260731_7kgf_sam1", row, node, "BAD" if index % 4 else "GOOD"])
    specs = (
        profile.DatasetSpec(
            "elec", str(electrical), "electrical_gt", "electrical_e5", title="Electrical"
        ),
        profile.DatasetSpec(
            "copy", str(writeback), "other", "electrical_e5", title="Electrical copy"
        ),
        profile.DatasetSpec("human", str(optical), "optical_human", "optical_3", title="Human"),
    )
    assert registry.scheme("electrical_e5")
    return profile.ReportProfile(specs, ())


def test_profile_workbook_embeds_one_curve_grid_per_distinct_sweep_table(
    tmp_path: Path, sweep_profile: profile.ReportProfile
) -> None:
    registry = schemes.load_registry(user_path="/nonexistent")
    result = profile.evaluate_profile(sweep_profile, registry=registry)
    assert not result.blocked, result.errors
    assert [spec.id for spec, _table in sweep_sources(result, result.samples[0])] == ["elec"]
    assert sweep_map_count(result) == 1
    with_codes, color_only = generate_workbook_pair(tmp_path / "report.xlsx", result)
    for path in (with_codes, color_only):
        wb = load_workbook(path)
        verify_profile_workbook(wb)
        assert wb["README"]["A8"].value == "Measurement sweep maps"
        assert wb["README"]["B8"].value == (
            "1 transfer-curve grid(s) embedded on the sample map sheets"
        )
        maps = wb["S01"]
        assert len(maps._images) == 1
        # Block titles all carry an em dash; the notes under the grid image do not.
        headings = [
            maps.cell(row, 3).value
            for row in range(1, maps.max_row + 1)
            if isinstance(maps.cell(row, 3).value, str) and " — " in maps.cell(row, 3).value
        ]
        assert headings[0] == "Electrical — raw labels"
        assert headings[1] == "Electrical — measured transfer curves (988/988 traces)"
        assert headings[2] == "Electrical copy — raw labels"
        assert headings[3] == "Human — raw labels"
        note_rows = [
            maps.cell(row, 3).value
            for row in range(1, maps.max_row + 1)
            if isinstance(maps.cell(row, 3).value, str)
            and maps.cell(row, 3).value.startswith("Curve colour")
        ]
        assert len(note_rows) == 1 and "Electrical label" in note_rows[0]
        image = maps._images[0]
        assert image.anchor._from.col == 2  # column C
        assert (image.width, image.height) == (26 + 38 * 32 + 4, 14 + 26 * 24 + 14)
        wb.close()
    embedded = (
        profile.ReportProfile.load(with_codes.with_suffix(".profile.json"))
        if (with_codes.with_suffix(".profile.json").exists())
        else sweep_profile
    )
    assert embedded.options.sweep_maps is True

    # Switching the option off keeps the label maps and drops the image and its record.
    off = profile.ReportProfile(
        sweep_profile.datasets, (), options=profile.ProfileOptions(sweep_maps=False)
    )
    result_off = profile.evaluate_profile(off, registry=registry)
    assert sweep_map_count(result_off) == 0
    with_codes_off, _ = generate_workbook_pair(tmp_path / "off.xlsx", result_off)
    wb = load_workbook(with_codes_off)
    verify_profile_workbook(wb)
    assert wb["README"]["B8"].value == "off"
    assert len(wb["S01"]._images) == 0
    assert profile.ProfileOptions.from_dict(off.options.to_dict()) == off.options
    assert profile.ProfileOptions.from_dict({"title": "x"}).sweep_maps is True
    wb.close()


def test_verify_rejects_a_missing_or_extra_curve_grid(
    tmp_path: Path, sweep_profile: profile.ReportProfile
) -> None:
    registry = schemes.load_registry(user_path="/nonexistent")
    result = profile.evaluate_profile(sweep_profile, registry=registry)
    with_codes, _ = generate_workbook_pair(tmp_path / "report.xlsx", result)
    wb = load_workbook(with_codes)
    wb["README"]["B8"] = "2 transfer-curve grid(s) embedded on the sample map sheets"
    with pytest.raises(ValueError, match="records 2 transfer-curve grid"):
        verify_profile_workbook(wb)
    wb["README"]["B8"] = "off"
    with pytest.raises(ValueError, match="carry 1 image"):
        verify_profile_workbook(wb)
    wb.close()
