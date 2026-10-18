"""Offscreen, synchronous (test_mode) walk through the five-step wizard."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QComboBox, QDialog

from r2r_evaluation_report import profile, schemes, wizard_backend
from r2r_evaluation_report.core import default_sample_selection as core_default
from r2r_evaluation_report.profile_workbook import verify_profile_workbook
from r2r_evaluation_report.wizard import NO_MEMBER, ComparisonDialog, WizardWindow

E5 = ("E-Normal", "E-NoActive", "E-Open", "E-Short", "E-Invalid")
GRID = [(r, n) for r in range(1, 27) for n in range(1, 39)]


def _write(path: Path, name: str, column: str, values) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["Name", "Row", "Node", column])
        writer.writeheader()
        for index, (row, node) in enumerate(GRID):
            writer.writerow({"Name": name, "Row": row, "Node": node, column: values(index)})
    return path


@pytest.fixture
def files(tmp_path: Path) -> dict[str, Path]:
    def elec(index):
        return E5[index % 5]

    def optical(index):
        return "BAD" if elec(index) in {"E-Invalid", "E-Short"} and index % 3 else "GOOD"

    def ml(index):
        return {"E-Normal": "Normal", "E-NoActive": "Normal", "E-Open": "Open"}.get(
            elec(index), "Short"
        )

    return {
        "elec": _write(tmp_path / "elec.csv", "20260731_7kgf_sam1", "Status", elec),
        "human": _write(tmp_path / "human.csv", "20260731_sam1_slices", "label", optical),
        "ml": _write(tmp_path / "ml.csv", "20260731_7kgf_sam1", "prediction", ml),
    }


def test_backend_inspects_loads_and_validates(files: dict[str, Path], tmp_path: Path) -> None:
    info = wizard_backend.inspect_source(files["elec"])
    assert info.worksheets == ("CSV",) and not info.is_inspector
    specs = [
        profile.DatasetSpec("a", str(files["elec"]), "electrical_gt"),
        profile.DatasetSpec("b", str(files["human"]), "optical_human"),
    ]
    registry = schemes.load_registry(user_path="/nonexistent")
    load = wizard_backend.load_and_propose(specs, registry)
    assert load.suggested_schemes == {"a": "electrical_e5", "b": "optical_3"}
    rows = wizard_backend.alignment_rows(specs, load.datasets, load)
    assert rows[0]["members"]["b"] == "20260731_sam1_slices" and not rows[0]["auto"]["b"]
    bad = profile.ComparisonSpec(
        "x", "reference", "a", "b", mapping_a={"E-Normal": "Normal"}, categories=("Normal",)
    )
    problems = wizard_backend.validate_comparison(bad, E5, ("GOOD", "BAD"))
    assert any("A면 라벨에 매핑이 없습니다" in item for item in problems)
    assert any("B면 라벨에 매핑이 없습니다" in item for item in problems)
    with pytest.raises(wizard_backend.DataContractError):
        wizard_backend.load_and_propose([], registry)


def test_wizard_walks_all_steps_and_generates(qtbot, files: dict[str, Path], tmp_path) -> None:
    registry = schemes.load_registry(user_path="/nonexistent")
    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    window.show()
    assert not window.next_button.isEnabled()
    elec = window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    human = window.add_dataset(str(files["human"]), role="optical_human", title="Human")
    ml = window.add_dataset(str(files["ml"]), role="electrical_ml", title="ML")
    assert [spec.id for spec in window.dataset_specs()] == [elec, human, ml]
    window.run_preflight()
    assert window.load is not None and window.alignment_box.isVisible()
    # The exact-title ML sheet is auto-selected; the human sheet needs confirmation.
    human_box = window.alignment_table.cellWidget(0, 1)
    ml_box = window.alignment_table.cellWidget(0, 2)
    assert isinstance(human_box, QComboBox) and human_box.currentText() == NO_MEMBER
    assert ml_box.currentText() == "20260731_7kgf_sam1"
    window._confirm_all()
    assert human_box.currentText() == "20260731_sam1_slices"
    assert len(window.confirmed_alignments()) == 2
    assert window.next_button.isEnabled()
    window._next()
    assert window.stack.currentIndex() == 1
    schemes_chosen = {spec.id: spec.scheme for spec in window.dataset_specs()}
    assert schemes_chosen == {elec: "electrical_e5", human: "optical_3", ml: "ml_3class"}
    # A wrong scheme blocks the step; restoring it unblocks.
    window.set_scheme(human, "legacy_electrical")
    window._refresh_scheme_table()
    assert not window.next_button.isEnabled()
    window.set_scheme(human, "optical_3")
    window._refresh_scheme_table()
    assert window.next_button.isEnabled()
    window._next()
    assert window.stack.currentIndex() == 2
    e5 = registry.scheme("electrical_e5").labels
    window.add_comparison(
        profile.ComparisonSpec(
            "self",
            "reference",
            elec,
            ml,
            "Electrical vs ML",
            mapping_a=registry.preset("e5_to_3class").resolved_rules(e5),
            mapping_b={label: label for label in ("Normal", "Open", "Short")},
            categories=("Normal", "Open", "Short"),
        )
    )
    window.add_comparison(
        profile.ComparisonSpec(
            "cross",
            "association",
            elec,
            human,
            "Electrical vs human",
            cells=(profile.CellSpec("E-Invalid x BAD", ("E-Invalid",), ("BAD",)),),
        )
    )
    assert window.comparison_table.rowCount() == 2
    assert window.comparison_table.item(0, 5).text() == "확인"
    window._next()
    assert window.stack.currentIndex() == 3
    assert not window.next_button.isEnabled()
    output = tmp_path / "out" / "wizard.xlsx"
    output.parent.mkdir()
    window.output_edit.setText(str(output))
    window.title_edit.setText("wizard run")
    window._refresh_next()
    assert window.next_button.isEnabled()
    window._next()
    assert window.stack.currentIndex() == 4
    assert "Electrical vs ML" in window.summary_view.toPlainText()
    window.run_generate()
    assert window._last_outputs is not None, window.status_label.text()
    with_codes, color_only = window._last_outputs
    assert with_codes == output and color_only is not None and color_only.exists()
    assert output.with_suffix(".profile.json").exists()
    wb = load_workbook(with_codes)
    verify_profile_workbook(wb)
    assert wb.sheetnames[-3:] == ["C01", "C02", "Overall Summary"]
    saved = profile.ReportProfile.load(output.with_suffix(".profile.json"))
    assert [c.id for c in saved.comparisons] == ["self", "cross"]
    assert len(saved.alignments) == 1 and set(saved.alignments[0].members) == {elec, human, ml}

    # Reloading the profile into a fresh wizard restores datasets, comparisons and
    # (after preflight) the confirmed alignments.
    again = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(again)
    again.apply_profile(saved)
    assert [c.id for c in again.comparisons] == ["self", "cross"]
    again.run_preflight()
    assert len(again.confirmed_alignments()) == 2
    assert again.build_profile().alignments == saved.alignments

    # Editing step 1 invalidates the load and later outputs.
    window._go(0)
    window._remove_dataset()  # nothing selected: no-op
    window.dataset_table.selectRow(2)
    window._remove_dataset()
    assert window.load is None and window._last_outputs is None
    assert [c.id for c in window.comparisons] == ["cross"]


def test_comparison_dialog_builds_reference_and_association(qtbot, files, tmp_path) -> None:
    registry = schemes.load_registry(user_path="/nonexistent")
    specs = [
        profile.DatasetSpec("a", str(files["elec"]), "electrical_gt", "electrical_e5", title="E"),
        profile.DatasetSpec("b", str(files["ml"]), "electrical_ml", "ml_3class", title="M"),
        profile.DatasetSpec("c", str(files["human"]), "optical_human", "optical_3", title="H"),
    ]
    labels = {"a": E5, "b": ("Normal", "Open", "Short"), "c": ("GOOD", "BAD")}
    dialog = ComparisonDialog(None, specs, labels, registry)
    qtbot.addWidget(dialog)
    dialog.side_widgets["A"]["preset"].setCurrentIndex(
        dialog.side_widgets["A"]["preset"].findData("e5_to_3class")
    )
    built = dialog.build()
    assert built.kind == "reference" and built.a == "a" and built.b == "b"
    assert built.mapping_a["E-Invalid"] == "Exclude" and built.mapping_b["Open"] == "Open"
    assert built.categories == ("Normal", "Open", "Short") and built.strict_macro
    assert built.preset_a == "e5_to_3class" and built.preset_b is None
    assert not wizard_backend.validate_comparison(built, labels["a"], labels["b"])
    dialog._accept()
    assert dialog.result_spec == built

    other = ComparisonDialog(None, specs, labels, registry, used_ids={"cmp1"})
    qtbot.addWidget(other)
    other.kind_box.setCurrentIndex(1)
    other.b_box.setCurrentIndex(2)
    other._add_cell("inv x bad", "E-Invalid", "BAD")
    assoc = other.build()
    assert assoc.id == "cmp2" and assoc.kind == "association" and assoc.b == "c"
    assert assoc.mapping_a == {} and assoc.mapping_b == {} and assoc.categories == ()
    assert assoc.cells == (profile.CellSpec("inv x bad", ("E-Invalid",), ("BAD",)),)
    reopened = ComparisonDialog(None, specs, labels, registry, existing=assoc)
    qtbot.addWidget(reopened)
    assert reopened.build() == assoc


def test_wizard_asks_for_sample_blocks_and_persists_choice(qtbot, tmp_path) -> None:
    from r2r_evaluation_report.wizard import SampleSelectionDialog
    from tests.test_profile import _merged_converter_csv

    source = _merged_converter_csv(tmp_path / "merged.csv")
    registry = schemes.load_registry(user_path="/nonexistent")
    specs = [profile.DatasetSpec("d1", str(source), "electrical_gt")]
    pending = wizard_backend.pending_sample_selections(specs)
    assert set(pending) == {"d1"} and len(pending["d1"]) == 7
    assert (
        wizard_backend.pending_sample_selections(
            [profile.DatasetSpec("d1", str(source), samples={"CSV::x_sam7::0": "x_sam7"})]
        )
        == {}
    )

    dialog = SampleSelectionDialog(None, "Electrical", pending["d1"])
    qtbot.addWidget(dialog)
    # Defaults: every complete block checked, the partial one disabled, duplicates flagged.
    assert dialog.current_selection() == core_default(pending["d1"])
    assert not dialog.table.cellWidget(5, 0).isEnabled()
    assert dialog.table.item(4, 8).text().startswith("동일 내용")
    assert dialog.table.isRowHidden(2) and not dialog.table.isRowHidden(0)  # sam7 has no issue
    dialog._bulk(first_only=True)
    assert dialog.current_selection() == {
        "CSV::x_sam5::0": "x_sam5",
        "CSV::x_sam1::0": "x_sam1",
        "CSV::x_sam7::0": "x_sam7",
        "CSV::x_sam19::1": "x_sam19",
    }
    dialog.table.item(0, 1).setText("x_sam1")  # clash with another title -> blocked
    dialog._accept()
    assert dialog.result() != QDialog.Accepted and "중복" in dialog.problems_label.text()
    dialog.table.item(0, 1).setText("x_sam5 (first)")
    dialog._accept()
    assert dialog.result() == QDialog.Accepted
    assert dialog.selection["CSV::x_sam5::0"] == "x_sam5 (first)"

    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    dataset_id = window.add_dataset(str(source), role="electrical_gt", title="Electrical")
    window.run_preflight()  # test_mode picks the default selection automatically
    assert window.load is not None
    assert len(window.load.datasets[dataset_id].sheets) == 6
    assert window.dataset_specs()[0].samples == core_default(pending["d1"])
    window.set_samples(dataset_id, dialog.selection)
    assert window.load is None  # selection change invalidates the load
    window.run_preflight()
    titles = [sheet.title for sheet in window.load.datasets[dataset_id].sheets]
    assert titles == ["x_sam5 (first)", "x_sam1", "x_sam7", "x_sam19"]  # first sam5 block kept
    saved = window.build_profile()
    assert saved.datasets[0].samples == dialog.selection


def _inspector_matrix(path: Path, name: str) -> Path:
    """Inspector matrix export with one cloud and one LM Studio (local) model column."""
    rows = []
    for index, (row, node) in enumerate(GRID):
        verdict = ("GOOD", "BAD", "OPEN")[index % 3]
        rows.append(
            {
                "image_path": f"C:/slices/{name}_rgb_{row}_{node}.png",
                "image_sha256": "abc",
                "agreement": "full",
                "gemini:gemini-2.5-pro": verdict,
                "gemini:gemini-2.5-pro actual_model_id": "gemini-2.5-pro",
                "gemini:gemini-2.5-pro reason": "",
                "lmstudio:qwen2.5-vl-7b-instruct": verdict if index % 7 else "GOOD",
                "lmstudio:qwen2.5-vl-7b-instruct actual_model_id": "qwen2.5-vl-7b-instruct@q6_k",
                "lmstudio:qwen2.5-vl-7b-instruct reason": "",
            }
        )
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return path


def test_local_llm_is_a_distinct_comparison_group(qtbot, files, tmp_path: Path) -> None:
    """An LM Studio column is pre-sorted into the Local LLM group and reported as such."""
    export = _inspector_matrix(tmp_path / "inspector.csv", "20260731_7kgf_sam1")
    assert profile.suggest_role(True, "lmstudio:qwen2.5-vl-7b-instruct") == "optical_vlm_local"
    assert profile.suggest_role(True, "gemini:gemini-2.5-pro") == "optical_vlm"
    assert profile.suggest_role(False, None) is None
    assert profile.is_local_llm_model("LMStudio:x") and not profile.is_local_llm_model("verdict")

    registry = schemes.load_registry(user_path="/nonexistent")
    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    elec = window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    cloud = window.add_dataset(str(export))  # first model column, no role given
    local = window.add_dataset(str(export), model_column="lmstudio:qwen2.5-vl-7b-instruct")
    specs = {spec.id: spec for spec in window.dataset_specs()}
    assert specs[cloud].role == "optical_vlm"
    assert specs[cloud].title == "inspector [gemini:gemini-2.5-pro]"
    assert specs[local].role == "optical_vlm_local"
    assert specs[local].title == "inspector [lmstudio:qwen2.5-vl-7b-instruct]"
    assert specs[local].provider == "lmstudio" and specs[local].is_local_llm
    assert not specs[cloud].is_local_llm
    assert specs[local].group_label.startswith("광학 Local LLM 라벨")

    # Switching the model column of a row re-suggests the group and the default title;
    # an operator-edited title is left alone.
    row = [spec.id for spec in window.dataset_specs()].index(cloud)
    model_box = window.dataset_table.cellWidget(row, 3)
    model_box.setCurrentText("lmstudio:qwen2.5-vl-7b-instruct")
    changed = {spec.id: spec for spec in window.dataset_specs()}[cloud]
    assert changed.role == "optical_vlm_local"
    assert changed.title == "inspector [lmstudio:qwen2.5-vl-7b-instruct]"
    window.dataset_table.item(row, 0).setText("My cloud run")
    model_box.setCurrentText("gemini:gemini-2.5-pro")
    changed = {spec.id: spec for spec in window.dataset_specs()}[cloud]
    assert changed.role == "optical_vlm" and changed.title == "My cloud run"

    window.run_preflight()
    assert window.load is not None
    window._confirm_all()
    assert len(window.confirmed_alignments()) == 2
    window._next()
    assert {spec.id: spec.scheme for spec in window.dataset_specs()}[local] == "optical_3"
    window._next()
    window.add_comparison(
        profile.ComparisonSpec("agree", "association", cloud, local, "Cloud vs Local LLM")
    )
    window.add_comparison(
        profile.ComparisonSpec(
            "elec",
            "association",
            elec,
            local,
            "Electrical vs Local LLM",
            cells=(profile.CellSpec("E-Short x BAD", ("E-Short",), ("BAD",)),),
        )
    )
    window._next()
    output = tmp_path / "local.xlsx"
    window.output_edit.setText(str(output))
    window._refresh_next()
    window._next()
    window.run_generate()
    assert window._last_outputs is not None, window.status_label.text()

    wb = load_workbook(output)
    verify_profile_workbook(wb)
    readme = [[c.value for c in r] for r in wb["README"].iter_rows()]
    header = next(r for r in readme if r[0] == "Id" and "Comparison group" in r)
    group_col, provider_col = header.index("Comparison group"), header.index("Provider")
    local_rows = [r for r in readme if r[0] == local]
    assert local_rows and local_rows[0][group_col].startswith("광학 Local LLM 라벨")
    assert local_rows[0][provider_col] == "lmstudio (local)"
    cloud_rows = [r for r in readme if r[0] == cloud]
    assert cloud_rows[0][provider_col] == "gemini (cloud)"
    summary = [[c.value for c in r] for r in wb["Overall Summary"].iter_rows()]
    assert any(r[0] == "Comparison groups (dataset roles)" for r in summary)
    groups = [r for r in summary if r and r[2] in ("lmstudio (local)", "gemini (cloud)")]
    assert {r[2] for r in groups} == {"lmstudio (local)", "gemini (cloud)"}
    assert any(isinstance(r[0], str) and r[0].startswith("Local LLM group:") for r in summary)
    saved = profile.ReportProfile.load(output.with_suffix(".profile.json"))
    assert saved.dataset(local).role == "optical_vlm_local"


def test_threaded_preflight_chains_inspection_and_load(qtbot, files, tmp_path: Path) -> None:
    """Real (non test_mode) preflight: the load job must start after the inspect job.

    Regression: the inspect job's continuation started the load job while the
    finished worker thread was still referenced, so _run refused it and the
    wizard sat at '이전 작업이 끝나기를 기다리세요.' with no datasets loaded; and the
    model column of Inspector datasets was dropped because the (disabled-while-busy)
    combo reported isEnabled() == False.
    """
    export = _inspector_matrix(tmp_path / "inspector.csv", "20260731_7kgf_sam1")
    registry = schemes.load_registry(user_path="/nonexistent")
    window = WizardWindow(test_mode=False, registry=registry)
    qtbot.addWidget(window)
    window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    local = window.add_dataset(str(export), model_column="lmstudio:qwen2.5-vl-7b-instruct")
    window.run_preflight()
    qtbot.waitUntil(lambda: window.load is not None, timeout=120_000)
    qtbot.waitUntil(lambda: window._thread is None, timeout=10_000)
    assert "2개 데이터셋" in window.status_label.text()
    assert window.load.suggested_schemes[local] == "optical_3"
    assert window.dataset_specs()[1].model_column == "lmstudio:qwen2.5-vl-7b-instruct"
    assert window.preflight_button.isEnabled()

    # ...and the threaded generate path delivers its result the same way.
    window._confirm_all()
    window._next()
    window._next()
    window.add_comparison(
        profile.ComparisonSpec("x", "association", window.dataset_specs()[0].id, local, "E x L")
    )
    window._next()
    output = tmp_path / "threaded.xlsx"
    window.output_edit.setText(str(output))
    window._refresh_next()
    window._next()
    window.run_generate()
    qtbot.waitUntil(lambda: window._last_outputs is not None, timeout=180_000)
    qtbot.waitUntil(lambda: window._thread is None, timeout=10_000)
    assert output.exists() and output.with_suffix(".profile.json").exists()


def _ml_export(path: Path, name: str, classes: tuple[str, str, str], verdicts) -> Path:
    """R2R-Machine-Learning predictions.csv with one ``prob_<class>`` column per class."""
    rows = []
    for index, (row, node) in enumerate(GRID):
        verdict = verdicts(index)
        rows.append(
            {
                "image_path": f"C:/slices/{name}_rgb_{row}_{node}.png",
                "name": name,
                "row": row,
                "node": node,
                "sha256": "abc",
                "run_date": "2026-10-17",
                "source_id": "slices",
                "prediction": verdict,
                "confidence": 0.9,
                "review_required": "false",
                **{f"prob_{label}": 0.9 if label == verdict else 0.05 for label in classes},
            }
        )
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return path


def test_visual_label_ml_export_is_its_own_comparison_group(qtbot, files, tmp_path: Path) -> None:
    """A Good/Bad/Open prediction set is pre-sorted into the visual-label ML group and
    compared against the human labels on one GOOD/BAD/OPEN axis."""
    visual_export = _ml_export(
        tmp_path / "predictions.csv",
        "20260731_sam1_slices",
        ("Good", "Bad", "Open"),
        lambda index: ("Good", "Bad", "Open")[index % 3] if index % 7 else "Good",
    )
    electrical_export = _ml_export(
        tmp_path / "electrical.csv",
        "20260731_7kgf_sam1",
        ("Normal", "Open", "Short"),
        lambda index: ("Normal", "Open", "Short")[index % 3],
    )
    info = wizard_backend.inspect_source(visual_export)
    assert info.prediction_classes == ("Good", "Bad", "Open")
    assert info.is_ml_prediction and not info.is_inspector
    assert wizard_backend.inspect_source(electrical_export).prediction_classes == (
        "Normal",
        "Open",
        "Short",
    )
    assert wizard_backend.inspect_source(files["ml"]).prediction_classes == ()
    # The same export saved as .xlsx is recognised through its header row.
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    with visual_export.open(encoding="utf-8-sig", newline="") as handle:
        for record in csv.reader(handle):
            sheet.append(record)
    xlsx = tmp_path / "predictions.xlsx"
    workbook.save(xlsx)
    assert wizard_backend.inspect_source(xlsx).prediction_classes == ("Good", "Bad", "Open")

    registry = schemes.load_registry(user_path="/nonexistent")
    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    elec = window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    human = window.add_dataset(str(files["human"]), role="optical_human", title="Human")
    visual = window.add_dataset(str(visual_export))  # no role given
    electrical_ml = window.add_dataset(str(electrical_export))
    specs = {spec.id: spec for spec in window.dataset_specs()}
    assert specs[visual].role == "optical_ml" and specs[visual].title == "predictions"
    assert specs[electrical_ml].role == "electrical_ml"
    assert specs[visual].group_label == profile.role_title("optical_ml")
    window.dataset_table.selectRow(3)
    window._remove_dataset()
    assert [spec.id for spec in window.dataset_specs()] == [elec, human, visual]

    window.run_preflight()
    assert window.load is not None
    assert window.load.suggested_schemes[visual] == "optical_3"
    assert window.load.raw_labels[visual] == ("Good", "Bad", "Open")
    window._confirm_all()
    assert len(window.confirmed_alignments()) == 2
    window._next()
    assert {spec.id: spec.scheme for spec in window.dataset_specs()}[visual] == "optical_3"
    window._next()

    # The comparison dialog offers the scheme spelling for the ML side, so the default
    # identity mapping human (A) vs visual ML (B) yields one GOOD/BAD/OPEN axis.
    specs = window.dataset_specs()
    ids = [spec.id for spec in specs]
    dialog = ComparisonDialog(window, specs, window._labels_by_dataset(), registry)
    qtbot.addWidget(dialog)
    dialog.a_box.setCurrentIndex(ids.index(human))
    dialog.b_box.setCurrentIndex(ids.index(visual))
    dialog.kind_box.setCurrentIndex(0)
    dialog.title_edit.setText("Human vs visual ML")
    reference = dialog.build()
    assert reference.kind == "reference"
    assert reference.mapping_b == {"GOOD": "GOOD", "BAD": "BAD", "OPEN": "OPEN"}
    assert reference.categories == ("GOOD", "BAD", "OPEN")
    assert not wizard_backend.validate_comparison(
        reference, window._labels_by_dataset()[human], window._labels_by_dataset()[visual]
    )
    window.add_comparison(reference)
    window.add_comparison(
        profile.ComparisonSpec(
            "elec_x_visual",
            "association",
            elec,
            visual,
            "Electrical vs visual ML",
            cells=(profile.CellSpec("E-Short x BAD", ("E-Short",), ("BAD",)),),
        )
    )
    assert window.comparison_table.item(0, 5).text() == "확인"
    window._next()
    output = tmp_path / "visual.xlsx"
    window.output_edit.setText(str(output))
    window._refresh_next()
    window._next()
    assert "Human vs visual ML" in window.summary_view.toPlainText()
    window.run_generate()
    assert window._last_outputs is not None, window.status_label.text()

    wb = load_workbook(output)
    verify_profile_workbook(wb)
    readme = [[c.value for c in r] for r in wb["README"].iter_rows()]
    header = next(r for r in readme if r[0] == "Id" and "Comparison group" in r)
    visual_row = next(r for r in readme if r[0] == visual)
    assert visual_row[header.index("Comparison group")] == profile.role_title("optical_ml")
    assert visual_row[header.index("Role key")] == "optical_ml"
    assert visual_row[header.index("Scheme")] == "optical_3"
    summary = [[c.value for c in r] for r in wb["Overall Summary"].iter_rows()]
    assert any(isinstance(r[0], str) and r[0].startswith("Visual-label ML group:") for r in summary)
    audit = [[c.value for c in r] for r in wb["Label_Audit"].iter_rows(min_row=3)]
    columns = audit[0]
    good = next(
        r for r in audit[1:] if r[0] == "predictions" and r[columns.index("Label")] == "GOOD"
    )
    assert good[columns.index("Observed spellings")] == "Good"
    saved = profile.ReportProfile.load(output.with_suffix(".profile.json"))
    assert saved.dataset(visual).role == "optical_ml"
    assert saved.comparisons[0].categories == ("GOOD", "BAD", "OPEN")


def test_partial_blocks_need_consent_in_dialog_and_flow_into_the_report(
    qtbot, files, tmp_path: Path
) -> None:
    from r2r_evaluation_report.wizard import SampleSelectionDialog
    from tests.test_profile import _partial_csv

    partial = _partial_csv(
        tmp_path / "human_partial.csv",
        "20260731_7kgf_sam1",
        "label",
        lambda index: ("GOOD", "BAD", "OPEN")[index % 3],
    )
    registry = schemes.load_registry(user_path="/nonexistent")
    candidates = wizard_backend.pending_sample_selections(
        [profile.DatasetSpec("d1", str(partial), "optical_human")]
    )["d1"]
    assert len(candidates) == 1 and not candidates[0].complete
    chosen = {candidates[0].key: "x"}
    assert wizard_backend.validate_selection(candidates, chosen)
    assert not wizard_backend.validate_selection(candidates, chosen, allow_partial=True)

    dialog = SampleSelectionDialog(None, "Human", candidates)
    qtbot.addWidget(dialog)
    box = dialog.table.cellWidget(0, 0)
    assert not box.isEnabled() and dialog.table.item(0, 8).text() == "불완전"
    assert "포함할 샘플이 없습니다" in dialog.problems_label.text()
    dialog._accept()
    assert dialog.result() != QDialog.Accepted
    dialog.allow_partial_box.setChecked(True)
    assert box.isEnabled() and dialog.table.item(0, 8).text() == "부분 (480/988)"
    dialog._bulk_including_partial()
    assert dialog.current_selection() == {candidates[0].key: "20260731_7kgf_sam1"}
    assert dialog.problems_label.text() == ""
    dialog.allow_partial_box.setChecked(False)  # withdrawing consent unticks the block
    assert not box.isChecked() and not box.isEnabled()
    dialog.allow_partial_box.setChecked(True)
    box.setChecked(True)
    dialog._accept()
    assert dialog.result() == QDialog.Accepted and dialog.allow_partial

    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    elec = window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    human = window.add_dataset(str(partial), role="optical_human", title="Human")
    window.set_samples(human, dialog.selection, dialog.allow_partial)
    assert {s.id: s.allow_partial for s in window.dataset_specs()} == {elec: False, human: True}
    window.run_preflight()
    assert window.load is not None, window.status_label.text()
    assert window.load.datasets[human].sheets[0].unique_coordinates == 480
    window._confirm_all()
    assert len(window.confirmed_alignments()) == 1
    window._next()
    window._next()
    window.add_comparison(
        profile.ComparisonSpec(
            "x",
            "association",
            elec,
            human,
            "E x H",
            cells=(profile.CellSpec("E-Short x BAD", ("E-Short",), ("BAD",)),),
        )
    )
    window._next()
    output = tmp_path / "partial.xlsx"
    window.output_edit.setText(str(output))
    window._refresh_next()
    window._next()
    assert "부분 블록 허용" in window.summary_view.toPlainText()
    window.run_generate()
    assert window._last_outputs is not None, window.status_label.text()
    saved = profile.ReportProfile.load(output.with_suffix(".profile.json"))
    assert saved.dataset(human).allow_partial and saved.dataset(human).samples == dialog.selection
    verify_profile_workbook(load_workbook(output))

    # Reloading the profile keeps the consent; the saved selection needs no new dialog.
    again = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(again)
    again.apply_profile(saved)
    assert again.dataset_specs()[1].allow_partial
    again.run_preflight()
    assert again.load is not None and len(again.confirmed_alignments()) == 1


def test_output_page_offers_the_transfer_curve_grid_option(qtbot, files, tmp_path) -> None:
    registry = schemes.load_registry(user_path="/nonexistent")
    window = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(window)
    window.add_dataset(str(files["elec"]), role="electrical_gt", title="Electrical")
    window.run_preflight()
    assert window.sweeps_check.isChecked() and window.sweeps_check.isEnabled()
    assert window.build_profile().options.sweep_maps is True
    window.sweeps_check.setChecked(False)
    assert window.build_profile().options.sweep_maps is False
    # The grid lives on the map sheets: without them the option is inert and greyed out.
    window.maps_check.setChecked(False)
    assert not window.sweeps_check.isEnabled()
    window.maps_check.setChecked(True)
    assert window.sweeps_check.isEnabled()
    window.output_edit.setText(str(tmp_path / "out.xlsx"))
    assert "transfer curve 격자 제외" in window._summary_text()
    saved = window.build_profile()
    assert profile.ReportProfile.from_dict(saved.to_dict()).options.sweep_maps is False
    again = WizardWindow(test_mode=True, registry=registry)
    qtbot.addWidget(again)
    again.apply_profile(saved)
    assert not again.sweeps_check.isChecked()
    again.sweeps_check.setChecked(True)
    again.output_edit.setText(str(tmp_path / "out.xlsx"))
    assert "transfer curve 격자 포함" in again._summary_text()
